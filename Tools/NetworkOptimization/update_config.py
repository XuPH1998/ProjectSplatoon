"""Append only the two authorized transport settings to the Luban source workbook."""
from copy import copy
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
PATH = ROOT / "Config/Luban/source/TbGlobal.xlsx"
FIELDS = (
    ("snapshotBytesPerSecond", "int", "全房间补同步载荷预算（字节/秒，必须大于0）", 262144),
    ("snapshotMaxInFlightRecords", "int", "每客户端未确认补同步记录上限（条，1到32）", 8),
)

def main():
    book = openpyxl.load_workbook(PATH)
    sheet = book["Global"]
    before = {(cell.coordinate): (cell.value, copy(cell._style)) for row in sheet for cell in row}
    headers = {cell.value: cell.column for cell in sheet[1]}
    for name, kind, description, value in FIELDS:
        if name in headers:
            assert [sheet.cell(row, headers[name]).value for row in range(1, 5)] == [name, kind, description, value]
            column = headers[name]
        else:
            column = sheet.max_column + 1
            for row, data in enumerate((name, kind, description, value), 1):
                cell = sheet.cell(row, column, data)
                cell._style = copy(sheet.cell(row, headers["chunksPerFrame"])._style)
        sheet.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 64
    book.save(PATH)
    saved = openpyxl.load_workbook(PATH)
    assert saved.sheetnames == book.sheetnames
    for coordinate, (value, style) in before.items():
        # openpyxl normalizes existing empty strings to empty cells on read.
        normalized = lambda item: None if item == "" else item
        assert normalized(saved["Global"][coordinate].value) == normalized(value), coordinate
        assert saved["Global"][coordinate]._style == style, coordinate
    for name, kind, description, value in FIELDS:
        column = next(cell.column for cell in saved["Global"][1] if cell.value == name)
        assert [saved["Global"].cell(row, column).value for row in range(1, 5)] == [name, kind, description, value]
    print("TbGlobal: added transport budget/window; existing values and cell styles preserved.")

if __name__ == "__main__":
    main()
