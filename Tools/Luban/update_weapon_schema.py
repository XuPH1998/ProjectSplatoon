#!/usr/bin/env python3
"""Validate the weapon/attachment Luban source schema without rewriting source workbooks.

This file used to be a destructive one-shot importer from an early reference sheet.
The live Excel files now contain hand-tuned camera shake, item metadata and attachment
data, so rebuilding them from that stale reference would discard valid columns.  Keep
this command as a fail-fast schema guard; edit the three source workbooks directly and
run ``Tools/Luban/gen_client.ps1`` to regenerate client outputs.
"""
from __future__ import annotations

import math
import sys
import xml.etree.ElementTree as ET
from pathlib import Path

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover - local tool dependency failure
    raise SystemExit("openpyxl is required to validate Luban workbooks") from exc


ROOT = Path(__file__).resolve().parents[2]
LUBAN = ROOT / "Assets" / "DesignData" / "Luban"
DEFINES = ROOT / "Tools" / "Luban" / "source" / "Defines" / "weapon.xml"

EXPECTED_WEAPON_IDS = set(range(1001, 1012))
EXPECTED_WEAPON_ITEM_IDS = set(range(3101, 3112))
EXPECTED_MELEE_WEAPON_IDS = {1012}
EXPECTED_MELEE_WEAPON_ITEM_IDS = {3112}
EXPECTED_MELEE_SKILL_PATHS = {
    1012: "Assets/GameResource/SharedCombat/Skills/SK_Player_Melee_1H_Attack.asset",
}
EXPECTED_ATTACHMENT_IDS = set(range(23001, 23026))
EXPECTED_ATTACHMENT_ITEM_IDS = set(range(3301, 3326))
EXPECTED_SUPPRESSOR_IDS = {23001, 23002, 23003}
EXPECTED_WEAPON_AUDIO_PROFILES = {
    1001: "Assets/GameResource/SharedCombat/Weapons/Audio/Profiles/SO_WeaponAudio_1001_AK_103.asset",
    1002: "Assets/GameResource/SharedCombat/Weapons/Audio/Profiles/SO_WeaponAudio_1002_AR_15.asset",
    1003: "Assets/GameResource/SharedCombat/Weapons/Audio/Profiles/SO_WeaponAudio_1003_MP5.asset",
    1004: "Assets/GameResource/SharedCombat/Weapons/Audio/Profiles/SO_WeaponAudio_1004_UZI.asset",
    1005: "Assets/GameResource/SharedCombat/Weapons/Audio/Profiles/SO_WeaponAudio_1005_PM.asset",
    1006: "Assets/GameResource/SharedCombat/Weapons/Audio/Profiles/SO_WeaponAudio_1006_Glock.asset",
    1007: "Assets/GameResource/SharedCombat/Weapons/Audio/Profiles/SO_WeaponAudio_1007_MP_155.asset",
    1008: "Assets/GameResource/SharedCombat/Weapons/Audio/Profiles/SO_WeaponAudio_1008_TOZ_66.asset",
    1009: "Assets/GameResource/SharedCombat/Weapons/Audio/Profiles/SO_WeaponAudio_1009_Desert_Eagle.asset",
    1010: "Assets/GameResource/SharedCombat/Weapons/Audio/Profiles/SO_WeaponAudio_1010_Mosin_Nagant.asset",
    1011: "Assets/GameResource/SharedCombat/Weapons/Audio/Profiles/SO_WeaponAudio_1011_Test_Weapon.asset",
}
EXPECTED_SLOT_NAMES = {
    "Grip",
    "Muzzle",
    "Stock",
    "Magazine",
    "Optic",
    "Tactical",
}
EXPECTED_STAT_NAMES = {
    "FireRate",
    "MagazineCapacity",
    "ReloadTime",
    "AdsTime",
    "MoveSpeedMul",
    "AdsMoveSpeedMul",
    "HipSpreadBase",
    "HipSpreadMax",
    "HipSpreadGrowth",
    "HipSpreadRecover",
    "VerticalRecoil",
    "HorizontalRecoil",
    "SoundMultiplier",
    "Damage",
    "CritRate",
    "MaxAimDistance",
    "AimCameraMaxOffset",
}
EXPECTED_MODIFIER_MODES = {"FlatAdd", "RateAdd"}


def require(condition: bool, message: str) -> None:
    if not condition:
        raise ValueError(message)


def load_sheet(path: Path):
    require(path.is_file(), f"missing workbook: {path}")
    # Artifact Tool exports do not always cache worksheet dimensions, so regular
    # loading is required here. This validator never saves the workbook.
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
    return workbook, workbook.active


def load_named_sheet(path: Path, sheet_name: str):
    require(path.is_file(), f"missing workbook: {path}")
    workbook = openpyxl.load_workbook(path, read_only=False, data_only=False)
    require(sheet_name in workbook.sheetnames,
            f"{path.name} missing sheet {sheet_name}")
    return workbook, workbook[sheet_name]


def header_map(sheet) -> dict[str, int]:
    return {
        str(sheet.cell(1, column).value): column
        for column in range(1, sheet.max_column + 1)
        if sheet.cell(1, column).value
    }


def data_rows(sheet, headers: dict[str, int]):
    id_column = headers["id"]
    for row in range(5, sheet.max_row + 1):
        if sheet.cell(row, id_column).value is not None:
            yield row


def enum_values(root: ET.Element, enum_name: str) -> set[str]:
    enum = root.find(f"./enum[@name='{enum_name}']")
    require(enum is not None, f"weapon.xml missing enum {enum_name}")
    return {entry.attrib["name"] for entry in enum.findall("var")}


def validate_defines() -> None:
    root = ET.parse(DEFINES).getroot()
    require(enum_values(root, "AttachmentSlotType") == EXPECTED_SLOT_NAMES,
            "AttachmentSlotType values do not match the six-slot MVP")
    require(enum_values(root, "AttachmentStatType") == EXPECTED_STAT_NAMES,
            "AttachmentStatType values do not match supported runtime stats")
    require(enum_values(root, "AttachmentModifierMode") == EXPECTED_MODIFIER_MODES,
            "AttachmentModifierMode must contain FlatAdd and RateAdd")

    weapon = root.find("./bean[@name='Weapon']")
    melee_weapon = root.find("./bean[@name='MeleeWeapon']")
    weapon_item_meta = root.find("./bean[@name='WeaponItemMeta']")
    attachment = root.find("./bean[@name='Attachment']")
    modifier = root.find("./bean[@name='AttachmentStatModifier']")
    require(weapon is not None and weapon.find("./var[@name='attachmentSlots']") is not None,
            "Weapon.attachmentSlots is missing")
    require("MeleeWeapon" in enum_values(root, "ItemType"),
            "ItemType.MeleeWeapon is missing")
    require(melee_weapon is not None, "MeleeWeapon bean is missing")
    require({entry.attrib["name"] for entry in melee_weapon.findall("var")} == {
                "id", "displayName", "critRate", "critMultiplier", "armorBreak",
                "attackRange", "targetHalfAngle", "moveSpeedMul", "weaponPrefabPath",
                "skillConfigPath",
            },
            "MeleeWeapon bean fields do not match the meleeWeapon sheet")
    require(weapon.find("./var[@name='defaultAmmoId']") is None,
            "Weapon.defaultAmmoId must remain removed")
    item = root.find("./bean[@name='Item']")
    require(item is not None and item.find("./var[@name='canSell']") is not None,
            "Item.canSell is missing")
    require(weapon.find("./var[@name='weaponAudioProfilePath']") is not None,
            "Weapon.weaponAudioProfilePath is missing")
    require(weapon_item_meta is not None,
            "WeaponItemMeta bean is missing")
    weapon_item_meta_fields = {entry.attrib["name"] for entry in weapon_item_meta.findall("var")}
    require(weapon_item_meta_fields == {"soundMultiplier"},
            "WeaponItemMeta must only contain soundMultiplier; weapon durability belongs to TbItem.maxDurability")
    require(attachment is not None, "Attachment bean is missing")
    attachment_fields = {entry.attrib["name"] for entry in attachment.findall("var")}
    require(attachment_fields == {
                "id", "name", "slotType", "isSuppressor", "compatibleWeaponIds", "modifiers"
            },
            "Attachment bean fields must match the current attachment table")
    require(modifier is not None, "AttachmentStatModifier bean is missing")


def validate_weapons() -> dict[int, set[str]]:
    workbook, sheet = load_named_sheet(LUBAN / "TbWeapon.xlsx", "firearmWeapon")
    try:
        headers = header_map(sheet)
        require("attachmentSlots#sep=," in headers,
                "TbWeapon.xlsx missing attachmentSlots#sep=,")
        require("weaponAudioProfilePath" in headers,
                "TbWeapon.xlsx missing weaponAudioProfilePath")
        require("defaultAmmoId" not in headers,
                "TbWeapon.xlsx must not contain defaultAmmoId")
        weapon_slots: dict[int, set[str]] = {}
        for row in data_rows(sheet, headers):
            weapon_id = int(sheet.cell(row, headers["id"]).value)
            require(weapon_id not in weapon_slots, f"duplicate weapon ID {weapon_id}")
            slots = str(sheet.cell(row, headers["attachmentSlots#sep=,"]).value or "")
            parsed = [slot.strip() for slot in slots.split(",") if slot.strip()]
            require(len(parsed) == len(set(parsed)), f"weapon {weapon_id} has duplicate slots")
            require(set(parsed) <= EXPECTED_SLOT_NAMES, f"weapon {weapon_id} has unknown slots")
            profile_path = str(sheet.cell(row, headers["weaponAudioProfilePath"]).value or "")
            require(profile_path == EXPECTED_WEAPON_AUDIO_PROFILES[weapon_id],
                    f"weapon {weapon_id} has unexpected weaponAudioProfilePath {profile_path!r}")
            weapon_slots[weapon_id] = set(parsed)
        require(set(weapon_slots) == EXPECTED_WEAPON_IDS,
                "TbWeapon.xlsx weapon IDs changed unexpectedly")
        return weapon_slots
    finally:
        workbook.close()


def validate_melee_weapons() -> set[int]:
    workbook, sheet = load_named_sheet(LUBAN / "TbWeapon.xlsx", "meleeWeapon")
    try:
        headers = header_map(sheet)
        required = {
            "id", "displayName", "critRate", "critMultiplier", "armorBreak",
            "attackRange", "targetHalfAngle", "moveSpeedMul", "weaponPrefabPath",
            "skillConfigPath",
        }
        require(required <= headers.keys(), "meleeWeapon sheet is missing required columns")
        ids: set[int] = set()
        for row in data_rows(sheet, headers):
            weapon_id = int(sheet.cell(row, headers["id"]).value)
            require(weapon_id not in ids, f"duplicate melee weapon ID {weapon_id}")
            require(float(sheet.cell(row, headers["attackRange"]).value) > 0,
                    f"melee weapon {weapon_id} must have positive attackRange")
            require(0 < float(sheet.cell(row, headers["targetHalfAngle"]).value) <= 180,
                    f"melee weapon {weapon_id} has invalid targetHalfAngle")
            require(0 < float(sheet.cell(row, headers["moveSpeedMul"]).value) <= 1,
                    f"melee weapon {weapon_id} has invalid moveSpeedMul")
            skill_path = str(sheet.cell(row, headers["skillConfigPath"]).value or "")
            require(skill_path == EXPECTED_MELEE_SKILL_PATHS[weapon_id],
                    f"melee weapon {weapon_id} has unexpected skillConfigPath {skill_path!r}")
            ids.add(weapon_id)
        require(ids == EXPECTED_MELEE_WEAPON_IDS,
                "meleeWeapon IDs changed unexpectedly")
        return ids
    finally:
        workbook.close()


def validate_melee_weapon_items(melee_weapon_ids: set[int]) -> None:
    workbook, sheet = load_sheet(LUBAN / "TbItem.xlsx")
    try:
        headers = header_map(sheet)
        item_ids: set[int] = set()
        referenced_ids: set[int] = set()
        for row in data_rows(sheet, headers):
            if sheet.cell(row, headers["itemType"]).value != "MeleeWeapon":
                continue
            item_id = int(sheet.cell(row, headers["id"]).value)
            weapon_id = int(sheet.cell(row, headers["itemRefId"]).value)
            require(weapon_id in melee_weapon_ids,
                    f"melee item {item_id} references unknown melee weapon {weapon_id}")
            require(int(sheet.cell(row, headers["maxStackCount"]).value) == 1,
                    f"melee item {item_id} must not stack")
            require(int(sheet.cell(row, headers["maxDurability"]).value) == 0,
                    f"melee item {item_id} must not use durability in the first release")
            item_ids.add(item_id)
            referenced_ids.add(weapon_id)
        require(item_ids == EXPECTED_MELEE_WEAPON_ITEM_IDS,
                "melee weapon item IDs changed unexpectedly")
        require(referenced_ids == melee_weapon_ids,
                "not every melee weapon config has exactly one item row")
    finally:
        workbook.close()


def validate_attachments(weapon_slots: dict[int, set[str]]) -> set[int]:
    workbook, sheet = load_sheet(LUBAN / "TbAttachment.xlsx")
    try:
        headers = header_map(sheet)
        required = {
            "id", "slotType", "isSuppressor", "compatibleWeaponIds#sep=,", "modifiers#sep=|"
        }
        require(required <= headers.keys(), "TbAttachment.xlsx is missing MVP columns")
        require("tier" not in headers, "TbAttachment.xlsx must not contain legacy tier column")
        ids: set[int] = set()
        for row in data_rows(sheet, headers):
            attachment_id = int(sheet.cell(row, headers["id"]).value)
            require(attachment_id not in ids, f"duplicate attachment ID {attachment_id}")
            slot = str(sheet.cell(row, headers["slotType"]).value)
            require(slot in EXPECTED_SLOT_NAMES, f"attachment {attachment_id} has unknown slot {slot}")
            raw_is_suppressor = sheet.cell(row, headers["isSuppressor"]).value
            is_suppressor = raw_is_suppressor in (True, 1, "1", "true", "True")
            require(is_suppressor == (attachment_id in EXPECTED_SUPPRESSOR_IDS),
                    f"attachment {attachment_id} has invalid isSuppressor value {raw_is_suppressor!r}")
            compatible = str(sheet.cell(row, headers["compatibleWeaponIds#sep=,"]).value or "")
            compatible_values = [int(value.strip()) for value in compatible.split(",") if value.strip()]
            compatible_ids = set(compatible_values)
            require(len(compatible_values) == len(compatible_ids),
                    f"attachment {attachment_id} has duplicate compatible weapon IDs")
            require(compatible_ids <= weapon_slots.keys(),
                    f"attachment {attachment_id} references unknown weapons")
            raw_modifiers = str(sheet.cell(row, headers["modifiers#sep=|"]).value or "")
            for raw_modifier in raw_modifiers.split("|"):
                if not raw_modifier.strip():
                    continue

                parts = [part.strip() for part in raw_modifier.split(";")]
                require(len(parts) == 3,
                        f"attachment {attachment_id} has malformed modifier {raw_modifier!r}")
                stat, mode, raw_value = parts
                require(stat in EXPECTED_STAT_NAMES,
                        f"attachment {attachment_id} has unknown modifier stat {stat}")
                require(mode in EXPECTED_MODIFIER_MODES,
                        f"attachment {attachment_id} has unknown modifier mode {mode}")
                try:
                    value = float(raw_value)
                except ValueError as exc:
                    raise ValueError(
                        f"attachment {attachment_id} has non-numeric modifier value {raw_value!r}"
                    ) from exc
                require(math.isfinite(value),
                        f"attachment {attachment_id} has non-finite modifier value {raw_value!r}")
            ids.add(attachment_id)
        require(ids == EXPECTED_ATTACHMENT_IDS,
                "TbAttachment.xlsx attachment IDs changed unexpectedly")
        return ids
    finally:
        workbook.close()


def validate_attachment_items(attachment_ids: set[int]) -> None:
    workbook, sheet = load_sheet(LUBAN / "TbItem.xlsx")
    try:
        headers = header_map(sheet)
        required = {
            "id",
            "itemType",
            "itemRefId",
            "value",
            "maxStackCount",
            "searchTime",
            "itemQuality",
            "itemIconPath",
            "canSell",
        }
        require(required <= headers.keys(), "TbItem.xlsx is missing required item columns")
        item_ids: set[int] = set()
        referenced_attachment_ids: set[int] = set()
        for row in data_rows(sheet, headers):
            item_id = int(sheet.cell(row, headers["id"]).value)
            raw_search_time = sheet.cell(row, headers["searchTime"]).value
            try:
                search_time = float(raw_search_time)
            except (TypeError, ValueError) as exc:
                raise ValueError(
                    f"item {item_id} has missing or non-numeric searchTime"
                ) from exc
            require(math.isfinite(search_time) and search_time >= 0,
                    f"item {item_id} has invalid searchTime {raw_search_time!r}")

            raw_quality = sheet.cell(row, headers["itemQuality"]).value
            try:
                item_quality = int(raw_quality)
            except (TypeError, ValueError) as exc:
                raise ValueError(
                    f"item {item_id} has missing or non-numeric itemQuality"
                ) from exc
            require(1 <= item_quality <= 5,
                    f"item {item_id} has itemQuality outside 1-5: {raw_quality!r}")

            item_icon_path = str(sheet.cell(row, headers["itemIconPath"]).value or "")
            require(item_icon_path.startswith(
                "Assets/GameResource/MainUI/Items/Icons/") and item_icon_path.endswith(".png"),
                    f"item {item_id} has invalid itemIconPath {item_icon_path!r}")

            raw_can_sell = sheet.cell(row, headers["canSell"]).value
            require(isinstance(raw_can_sell, bool),
                    f"item {item_id} has invalid canSell value {raw_can_sell!r}")
            if item_id == 3112:
                require(not raw_can_sell, "default sword item 3112 must not be sellable")
                require(int(sheet.cell(row, headers["value"]).value) == 0,
                        "default sword item 3112 must have zero value")

            if sheet.cell(row, headers["itemType"]).value != "Attachment":
                continue
            require(item_id not in item_ids, f"duplicate attachment item ID {item_id}")
            attachment_id = int(sheet.cell(row, headers["itemRefId"]).value)
            require(int(sheet.cell(row, headers["maxStackCount"]).value) == 1,
                    f"attachment item {item_id} must not stack")
            require(attachment_id in attachment_ids,
                    f"attachment item {item_id} references unknown attachment {attachment_id}")
            item_ids.add(item_id)
            referenced_attachment_ids.add(attachment_id)
        require(item_ids == EXPECTED_ATTACHMENT_ITEM_IDS,
                "TbItem.xlsx attachment item IDs changed unexpectedly")
        require(referenced_attachment_ids == attachment_ids,
                "not every attachment config has exactly one item row")
    finally:
        workbook.close()


def validate_weapon_items(weapon_ids: set[int]) -> None:
    workbook, sheet = load_sheet(LUBAN / "TbItem.xlsx")
    try:
        headers = header_map(sheet)
        required = {
            "id",
            "itemType",
            "itemRefId",
            "maxStackCount",
            "maxDurability",
            "durabilityCost",
        }
        require(required <= headers.keys(), "TbItem.xlsx is missing weapon durability columns")
        item_ids: set[int] = set()
        referenced_weapon_ids: set[int] = set()
        for row in data_rows(sheet, headers):
            if sheet.cell(row, headers["itemType"]).value != "Weapon":
                continue

            item_id = int(sheet.cell(row, headers["id"]).value)
            weapon_id = int(sheet.cell(row, headers["itemRefId"]).value)
            max_durability = int(sheet.cell(row, headers["maxDurability"]).value)
            require(item_id not in item_ids, f"duplicate weapon item ID {item_id}")
            require(weapon_id in weapon_ids,
                    f"weapon item {item_id} references unknown weapon {weapon_id}")
            require(int(sheet.cell(row, headers["maxStackCount"]).value) == 1,
                    f"weapon item {item_id} must not stack")
            require(max_durability > 0,
                    f"weapon item {item_id} must define a positive maxDurability")
            require(int(sheet.cell(row, headers["durabilityCost"]).value) == 0,
                    f"weapon item {item_id} must not use generic durabilityCost")
            item_ids.add(item_id)
            referenced_weapon_ids.add(weapon_id)

        require(item_ids == EXPECTED_WEAPON_ITEM_IDS,
                "TbItem.xlsx weapon item IDs changed unexpectedly")
        require(referenced_weapon_ids == weapon_ids,
                "not every weapon config has exactly one weapon item row")
    finally:
        workbook.close()


def main() -> int:
    try:
        validate_defines()
        weapon_slots = validate_weapons()
        validate_weapon_items(set(weapon_slots))
        melee_weapon_ids = validate_melee_weapons()
        validate_melee_weapon_items(melee_weapon_ids)
        attachment_ids = validate_attachments(weapon_slots)
        validate_attachment_items(attachment_ids)
    except (OSError, ValueError, ET.ParseError) as exc:
        print(f"Weapon attachment schema validation failed: {exc}", file=sys.stderr)
        return 1

    print("Weapon attachment Luban source schema is valid. Run Tools/Luban/gen_client.ps1 next.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
