import openpyxl

path = r"d:\Unity\ProjectVanguard\Assets\DesignData\Luban\TbWeapon.xlsx"
ADS_SPREAD_HEADERS = {
    "瞄准初始扩散（度）",
    "瞄准最大扩散（度）",
    "瞄准扩散增长（度/发）",
    "瞄准扩散恢复（度/秒）",
}


wb = openpyxl.load_workbook(path)
ws = wb.active

removed = []
for col in range(ws.max_column, 0, -1):
    if ws.cell(4, col).value in ADS_SPREAD_HEADERS:
        ws.delete_cols(col, 1)
        removed.append(col)

if removed:
    wb.save(path)

print(f"Removed ADS spread columns: {sorted(removed)}")
