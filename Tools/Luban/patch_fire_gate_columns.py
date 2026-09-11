from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
PATH = ROOT / "Assets" / "DesignData" / "Luban" / "TbWeapon.xlsx"
INSERT_AFTER_COMMENT = "瞄准角速度上限（度/秒）"

FIRE_GATE_FIELDS = (
    ("开火判定扇形半角（度）", 0),
    ("自动开火锥形检测半径（米）", 1),
    ("扇形内持续瞄准时间（秒）", 2),
    ("瞄准角速度上限（度/秒）", 3),
    ("锁定目标失效宽限时间（秒）", 4),
    ("战斗保持时锁定目标失效宽限时间（秒）", 5),
    ("射击后战斗保持时间（秒）", 6),
    ("战斗瞄准保持时间（秒）", 7),
)

# (coneHalfAngleDeg, radiusMeters, settleTimeSec, maxAimSpeedDegPerSec,
#  breakGraceSec, combatBreakGraceSec, combatSustainAfterShotSec, combatAimSustainSec)
FIRE_GATE_BY_WEAPON_ID = {
    1001: (7.0, 25.0, 0.10, 90.0, 0.20, 0.45, 0.55, 0.35),
    1002: (7.0, 25.0, 0.10, 88.0, 0.20, 0.45, 0.55, 0.35),
    1003: (9.0, 25.0, 0.10, 140.0, 0.20, 0.45, 0.55, 0.35),
    1004: (9.0, 25.0, 0.10, 150.0, 0.20, 0.45, 0.55, 0.35),
    1005: (5.0, 25.0, 0.10, 80.0, 0.20, 0.45, 0.55, 0.35),
    1006: (5.0, 25.0, 0.10, 85.0, 0.20, 0.45, 0.55, 0.35),
    1007: (10.0, 25.0, 0.10, 110.0, 0.20, 0.45, 0.55, 0.35),
    1008: (10.0, 25.0, 0.10, 115.0, 0.20, 0.45, 0.55, 0.35),
}


def unmerge_all(ws) -> None:
    for merged_range in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged_range))


def rebuild_header_merges(ws) -> None:
    col = 2
    while col <= ws.max_column:
        name = ws.cell(1, col).value
        type_name = ws.cell(2, col).value
        if name in ("combat", "ballistics", "handling", "spread", "recoil", "itemMeta"):
            end = col
            while end + 1 <= ws.max_column and ws.cell(1, end + 1).value in (None, ""):
                end += 1
            ws.cell(1, col).value = name
            ws.cell(2, col).value = type_name
            if end > col:
                ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=end)
                ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=end)
            col = end + 1
        else:
            col += 1


def header_comment_columns(ws) -> dict:
    return {ws.cell(4, c).value: c for c in range(1, ws.max_column + 1)}


wb = openpyxl.load_workbook(PATH)
ws = wb.active
comments = header_comment_columns(ws)

missing = [comment for comment, _ in FIRE_GATE_FIELDS if comment not in comments]
if missing:
    insert_after = comments.get(INSERT_AFTER_COMMENT)
    if insert_after is None:
        raise RuntimeError(f"Could not find insert point: {INSERT_AFTER_COMMENT}")

    unmerge_all(ws)
    ws.insert_cols(insert_after + 1, len(missing))
    for offset, comment in enumerate(missing):
        col = insert_after + 1 + offset
        ws.cell(3, col).value = "c"
        ws.cell(4, col).value = comment
    rebuild_header_merges(ws)
    comments = header_comment_columns(ws)

for row in range(5, ws.max_row + 1):
    weapon_id = ws.cell(row, 2).value
    values = FIRE_GATE_BY_WEAPON_ID.get(weapon_id)
    if values is None:
        continue
    for comment, value_index in FIRE_GATE_FIELDS:
        ws.cell(row, comments[comment]).value = values[value_index]

wb.save(PATH)
print("Patched fire gate columns in TbWeapon.xlsx")
