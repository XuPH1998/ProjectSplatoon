"""Add the approved foam settings while preserving the existing Luban workbook layout."""
from copy import copy
from pathlib import Path
import xml.etree.ElementTree as ET
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
FIELDS = {
    "Map": [
        ("foamCellSize", "float", "泡沫：高度采样间距（米）", 0.25),
        ("foamChunkSize", "float", "泡沫：分块边长（米）", 4),
        ("foamMaxHeight", "float", "泡沫：最大新增高度（米）", 3),
        ("foamCeilingGap", "float", "泡沫：顶部障碍间隙（米）", 0.1),
    ],
    "Global": [
        ("foamCommitRate", "int", "泡沫：权威地形提交频率（Hz，整除模拟频率）", 20),
        ("foamDissolveRatio", "float", "泡沫：敌方消融效率倍率", 1.5),
        ("foamSlopeDegrees", "float", "泡沫：自由坡面目标角度（度）", 35),
    ],
}

def main():
    source = ROOT / "Config/Luban/source"
    schema = source / "Defines/gameplay.xml"
    text = schema.read_text(encoding="utf-8-sig")
    for name, fields in FIELDS.items():
        path = source / f"Tb{name}.xlsx"
        book = openpyxl.load_workbook(path)
        sheet = book.active
        before = list(sheet.values)
        for key, kind, label, value in fields:
            names = [c.value for c in sheet[1]]
            col = names.index(key) + 1 if key in names else sheet.max_column + 1
            for row, item in enumerate((key, kind, label, value), 1):
                cell = sheet.cell(row, col)
                if col > len(before[0]):
                    cell._style = copy(sheet.cell(row, len(before[0]))._style)
                    cell.alignment = copy(sheet.cell(row, len(before[0])).alignment)
                cell.value = item
            sheet.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 30
            declaration = f'    <var name="{key}" type="{kind}" comment="{label}" />'
            if f'name="{key}"' not in text:
                start = text.index(f'<bean name="{name}Config"')
                end = text.index("  </bean>", start)
                text = text[:end] + declaration + "\n" + text[end:]
        assert [tuple(sheet.cell(r+1, c+1).value for c in range(len(row))) for r, row in enumerate(before)] == before
        book.save(path)
        assert list(openpyxl.load_workbook(path).active.values)[3][-len(fields):] == tuple(f[3] for f in fields)
        print(f"{path.name}: preserved existing cells, added {len(fields)} settings")
    ET.fromstring(text)
    schema.write_text(text, encoding="utf-8")

if __name__ == "__main__":
    main()
