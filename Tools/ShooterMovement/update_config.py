"""Idempotently migrate source workbooks and schema for shooter / ink movement."""
from pathlib import Path
from copy import copy
import re
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
ADDITIONS = {
    "Character": [
        ("shootMoveSpeed", "float", "射击移动速度（米/秒）", 3.6),
        ("moveAcceleration", "float", "人形加减速度（米/秒²）", 40),
        ("swimAcceleration", "float", "潜墨加减速度（米/秒²）", 50),
        ("wallSwimSpeed", "float", "墙游速度（米/秒）", 4),
        ("wallGraceSeconds", "float", "墙面接缝接触容错（秒）", .1),
        ("wallProbeDistance", "float", "角色表面探测距离（米）", .6),
        ("wallJumpSpeed", "float", "脱墙跳跃外推速度（米/秒）", 4),
        ("mantleSeconds", "float", "墨水形态翻越时长（秒）", .25),
        ("healthRecoverDelay", "float", "受伤后回血等待（秒）", 1),
        ("healthRecoverRate", "float", "普通回血（HP/秒）", 30),
        ("swimHealthRecoverRate", "float", "己方潜墨回血（HP/秒）", 60),
        ("enemyInkDamageRate", "float", "敌墨伤害（HP/秒）", 20),
        ("enemyInkHealthFloor", "float", "敌墨非致死生命下限（HP）", 60),
    ],
    "Weapon": [
        ("fireIntervalFrames", "int", "连发间隔（60Hz参考帧）", 6),
        ("startFrames", "int", "人形起手（60Hz参考帧）", 2),
        ("emergeStartFrames", "int", "出墨起手（60Hz参考帧）", 10),
        ("inkRecoverLockFrames", "int", "最后一发后回墨锁定（60Hz参考帧）", 20),
        ("damageMin", "float", "衰减后伤害（HP）", 18),
        ("damageReduceStartFrames", "int", "伤害衰减开始（飞行60Hz参考帧）", 8),
        ("damageReduceEndFrames", "int", "伤害衰减结束（飞行60Hz参考帧）", 40),
        ("straightFrames", "int", "直行阶段（60Hz参考帧）", 4),
        ("brakeFrames", "int", "减速过渡（60Hz参考帧）", 8),
        ("brakeSpeedMultiplier", "float", "减速后的速度比例", .66),
        ("jumpSpreadDegrees", "float", "跳跃散布半角（度）", 12),
        ("spreadRecoverFrames", "int", "落地散布恢复（60Hz参考帧）", 12),
        ("trailSpacing", "float", "沿途落墨间隔（米）", .7),
        ("trailRadiusMin", "float", "沿途落墨最小笔刷半径（米）", .464),
        ("trailRadiusMax", "float", "沿途落墨最大笔刷半径（米）", .696),
        ("trailMaxDrop", "float", "沿途落墨向下探测范围（米）", 2.4),
        ("effectiveRange", "float", "伤害弹最大前向射程（本项目米）", 10.4),
        ("paintRange", "float", "水平瞄准涂地射程目标（本项目米）", 13.6),
    ],
    "Global": [("simulationRate", "int", "玩法模拟和输入采样频率（Hz）", 60)],
}
VALUES = {
    "Weapon": {"fireRate": 10, "damage": 36, "shotInk": .92,
               "speedMin": 31, "speedMax": 31, "gravity": 9.8, "lifetime": 1.2,
               "paintRadiusMin": .65, "paintRadiusMax": .8, "paintHardness": .55},
    "Arena": {"layoutVersion": 4},
}

schema_path = ROOT / "Config/Luban/source/Defines/gameplay.xml"
schema = schema_path.read_text(encoding="utf-8")
for table in dict.fromkeys([*ADDITIONS, *VALUES]):
    path = ROOT / "Config/Luban/source" / f"Tb{table}.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    cols = {c.value: c.column for c in ws[1] if c.value}
    for key, kind, comment, value in ADDITIONS.get(table, []):
        if key not in cols:
            col = ws.max_column + 1
            for row, item in enumerate([key, kind, comment, value], 1):
                cell = ws.cell(row, col, item)
                cell._style = copy(ws.cell(row, col - 1)._style)
                cell.alignment = copy(ws.cell(row, col - 1).alignment)
            ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 28
            cols[key] = col
        ws.cell(4, cols[key], value)
    for key, value in VALUES.get(table, {}).items():
        ws.cell(4, cols[key], value)
    wb.save(path)
    check = openpyxl.load_workbook(path, read_only=True, data_only=True)
    assert all(check.worksheets[0].cell(4, cols[k]).value == v for k, v in VALUES.get(table, {}).items())
    check.close()
    pattern = rf'(<bean name="{table}Config"[^>]*>)(.*?)(  </bean>)'
    match = re.search(pattern, schema, re.S)
    body = match[2]
    for key, kind, comment, value in ADDITIONS.get(table, []):
        if f'name="{key}"' not in body:
            body += f'    <var name="{key}" type="{kind}" comment="{comment}"/>\n'
    schema = schema[:match.start()] + match[1] + body + match[3] + schema[match.end():]
    print(f"Migrated Tb{table}")
schema_path.write_text(schema, encoding="utf-8")
