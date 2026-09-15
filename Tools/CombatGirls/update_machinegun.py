"""Add ID 6 through the authoritative workbook; safe to rerun, never rebalance IDs 1-5.

openpyxl is the fallback for the desktop runtime without @oai/artifact-tool.
"""
from pathlib import Path
from copy import copy
import json
import hashlib
import urllib.request
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
REF = ROOT / 'Tools/ValidationData/MachineGun'
SCALE = 18 / 24.037
FIELDS = [
    ('splatlingMinChargeFrames', 'int', '旋转枪：最小有效蓄力帧数', 8),
    ('splatlingFirstChargeFrames', 'int', '旋转枪：第一圈蓄力帧数', 120),
    ('splatlingFirstShootFrames', 'int', '旋转枪：第一圈射击窗口帧数（含第0帧首弹）', 130),
    ('splatlingFullShootFrames', 'int', '旋转枪：满蓄射击窗口帧数（含末帧）', 260),
    ('splatlingSlowChargeMultiplier', 'float', '旋转枪：空中或缺墨蓄力耗时倍率（不叠乘）', 3),
    ('splatlingChargeMoveSpeed', 'float', '旋转枪：蓄力移动速度（米/秒）', 5 * .044 / .096),
    ('splatlingChargeJumpSpeed', 'float', '旋转枪：蓄力起跳速度（米/秒）', .06 * 60 * SCALE),
    ('splatlingPostFrames', 'int', '旋转枪：末弹后恢复帧数', 4),
    ('splatlingPitchSpread', 'float', '旋转枪：地面垂直散布半角（度）', 2),
    ('splatlingSpreadBias', 'float', '旋转枪：中心散布偏向（0到1）', .3),
    ('splatlingSpeedBias', 'float', '旋转枪：初速随机中心偏向（0到1）', .2),
    ('splatlingFootEvery', 'int', '旋转枪：脚下落墨间隔（发）', 8),
    ('splatlingTrailCount', 'int', '旋转枪：每颗沿途最大落墨数', 1),
    ('splatlingFootRadius', 'float', '旋转枪：脚下落墨半径（米）', 1.5456 * SCALE),
    ('splatlingPlayerRadius', 'float', '旋转枪：命中玩家扫掠半径（米）', .225 * SCALE),
]

def main():
    REF.mkdir(parents=True, exist_ok=True)
    path = ROOT / 'Config/Luban/source/TbHero.xlsx'
    wb = openpyxl.load_workbook(path); ws = wb['Hero']
    cols = {c.value: c.column for c in ws[1] if c.value}
    old = {int(ws.cell(r, cols['id']).value): {k: ws.cell(r,c).value for k,c in cols.items() if not k.startswith('##')} for r in range(4, ws.max_row+1) if ws.cell(r,cols['id']).value}
    baseline = REF / 'Heroes-Before.json'
    if not baseline.exists(): baseline.write_text(json.dumps([old[i] for i in range(1,6)], ensure_ascii=False, indent=2), 'utf-8')
    for key,kind,label,value in FIELDS:
        if key not in cols:
            cols[key] = ws.max_column + 1
            for r in range(1, ws.max_row+1):
                ws.cell(r,cols[key])._style = copy(ws.cell(r,cols['burstCount'])._style)
                ws.cell(r,cols[key]).value = (key,kind,label)[r-1] if r<=3 else 0
            ws.column_dimensions[openpyxl.utils.get_column_letter(cols[key])].width = 28
    row = next((r for r in range(4,ws.max_row+1) if ws.cell(r,cols['id']).value == 6), ws.max_row+1)
    if 6 not in old:
        for c in range(1,ws.max_column+1):
            ws.cell(row,c)._style = copy(ws.cell(4,c)._style)
            ws.cell(row,c).value = ws.cell(4,c).value
        values = dict(id=6,name='MachineGunGirl',displayName='消防栓旋转枪',characterPrefabAddress='Character/MachineGunGirl',weaponPrefabAddress='Weapon/MachineGun',
            moveSpeed=5*.088/.096,swimSpeed=7.2,shootMoveSpeed=5*.06/.096,fireMode=4,fireRate=15,shotInk=35/66,
            startFrames=0,emergeStartFrames=4,inkRecoverLockFrames=40,burstCount=1,burstRecoveryFrames=0,pelletCount=1,muzzleMode=0,semiBufferFrames=0,
            speedMin=2.26*60*SCALE,speedMax=2.54*60*SCALE,projectileGravity=.016*3600*SCALE,lifetime=3,
            collisionRadius=.2*SCALE,spreadDegrees=3,jumpSpreadDegrees=6,spreadRecoverFrames=45,straightFrames=8,brakeFrames=8,brakeSpeedMultiplier=.1,
            effectiveRange=27*SCALE,damage=40,damageMin=16,damageReduceStartFrames=11,damageReduceEndFrames=19,
            paintRadiusMin=1.92*SCALE,paintRadiusMax=2.1*SCALE,trailSpacing=23.5*SCALE,trailRadiusMin=1.288*SCALE,trailRadiusMax=1.288*SCALE,trailMaxDrop=10*SCALE,
            chargeFrames=150,chargeMinDamage=32,chargePartialMaxDamage=32,chargeMinInk=35/66,chargeMinRange=11*SCALE,chargeMinSpeed=1.05*60*SCALE,chargeMinSpread=3,chargeMinJumpSpread=6)
        values.update({k:v for k,_,_,v in FIELDS})
        for key,value in values.items(): ws.cell(row,cols[key]).value=value
    ws.cell(3,cols['fireMode']).value = '发射：0全自动／1三连发／2蓄力松开发射／3半自动／4旋转枪'
    ws.cell(3,cols['shotInk']).value = '发射：每次有效发射耗墨（霰弹整组；旋转枪每颗从预留量结算）'
    for i in range(1,6):
        r=next(r for r in range(4,ws.max_row+1) if ws.cell(r,cols['id']).value==i)
        for k,v in old[i].items(): assert ws.cell(r,cols[k]).value==v,(i,k)
    wb.save(path); wb.close()
    schema=ROOT/'Config/Luban/source/Defines/gameplay.xml'; text=schema.read_text('utf-8')
    needle='    <var name="semiBufferFrames"'
    at=text.index('\n',text.index(needle))+1
    additions=''.join(f'    <var name="{k}" type="{t}" comment="{label}" />\n' for k,t,label,_ in FIELDS if f'name="{k}"' not in text)
    text=text[:at]+additions+text[at:]
    text=text.replace('2蓄力松开发射／3半自动"','2蓄力松开发射／3半自动／4旋转枪"')
    schema.write_text(text,'utf-8')
    packs_path=ROOT/'Tools/CombatGirls/hero-packs.json'; packs=json.loads(packs_path.read_text())
    if not any(p['id']==6 for p in packs):
        packs.append(dict(id=6,pack='CombatGirls_MachineGun',name='MachineGunGirl',weapon='MachineGun',avatar='Humanoid_F',prefab='Prefab/MachineGun_Girl.prefab',scene='MachineGun_Girl_Scene.unity',clips=['MG_AimWalk_F','MG_AimWalk_B','MG_AimWalk_FL','MG_AimWalk_BR','MG_AimTurn_L90','MG_AimTurn_R90','MG_AimIdle','MG_Shoot','MG_Die_F','MG_Die_B']))
        packs_path.write_text(json.dumps(packs,indent=2)+'\n','utf-8')
    url='https://raw.githubusercontent.com/Leanny/splat3/7280ff9cde8bb1c5dcef46c700c326471584d2e6/data/parameter/1130/weapon/WeaponSpinnerHyper.game__GameParameterTable.json'
    raw=REF/'WeaponSpinnerHyper.1130.json'
    if not raw.exists(): raw.write_bytes(urllib.request.urlopen(url).read())
    (REF/'Reference.json').write_text(json.dumps(dict(version='11.3.0',source=url,sha256=hashlib.sha256(raw.read_bytes()).hexdigest(),spatialScale=SCALE,normalWalkingAnchor=5,swimAnchor=8,fullChargeShots=66,firstChargeShots=33),indent=2),'utf-8')
    print('ID 6 installed in source workbook; original hero values preserved.')

if __name__=='__main__': main()
