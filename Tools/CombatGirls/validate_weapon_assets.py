"""Read-only source workbook -> Luban -> weapon assets -> Addressables migration audit."""
from pathlib import Path
import json
import math
import re
import xml.etree.ElementTree as ET
import openpyxl
import argparse
from migrate_weapon_seconds import TIME_FIELDS, seconds_values

ROOT = Path(__file__).resolve().parents[2]

def asset_values(path):
    result = {}
    for key, value in re.findall(r'^  ([a-z]\w*): (.+)$', path.read_text('utf-8-sig'), re.M):
        if key.startswith('m_'): continue
        try: result[key] = json.loads(value)
        except json.JSONDecodeError: result[key] = value
    return result

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--hero', type=int, help='Audit one authored hero while retaining source/schema checks')
    parser.add_argument('--output', type=Path, help='Write this audit to a separate report directory')
    args = parser.parse_args()
    errors = []
    def check(value, message):
        if not value: errors.append(message)
    def equal(a, b):
        return math.isclose(a, b, rel_tol=1e-6, abs_tol=1e-6) if isinstance(a, (int, float)) and isinstance(b, (int, float)) else a == b
    book = openpyxl.load_workbook(ROOT/'Config/Luban/source/TbHero.xlsx', data_only=False)
    sheet = book['Hero']
    columns = {c.value: c.column for c in sheet[1] if c.value and not c.value.startswith('##')}
    rows = [{key: sheet.cell(r, col).value for key, col in columns.items()} for r in range(4, sheet.max_row+1) if sheet.cell(r, columns['id']).value]
    book.close()
    generated = json.loads((ROOT/'Assets/GameResource/Bootstrap/Config/Luban/tbhero.json').read_text('utf-8-sig'))
    baseline = json.loads((ROOT/'Tools/ValidationData/WeaponAssets/Migration-Baseline.json').read_text('utf-8'))
    pistol_tuning = json.loads((ROOT/'Tools/ValidationData/WeaponAssets/PistolGirl-SplashTuning.json').read_text('utf-8'))
    dualies_tuning = json.loads((ROOT/'Tools/ValidationData/WeaponAssets/DualPistolGirl-GloogaNormal.json').read_text('utf-8'))
    blaster_tuning = json.loads((ROOT/'Tools/ValidationData/WeaponAssets/RapidBlaster-Tuning.json').read_text('utf-8'))
    alignment = json.loads((ROOT/'Tools/ValidationData/WeaponAlignment/Tuning.json').read_text('utf-8'))
    def aligned(hero, values):
        values.update(next(r for r in json.loads((ROOT/'Tools/ValidationData/WeaponAlignment/Baseline/Values.json').read_text('utf-8')) if r['id']==hero))
        values.update({r['field']:r['after'] for r in alignment['fields'] if r['hero']==hero})
        if hero==1: values['weaponTypeName']='斯普拉射击枪'
        if hero==3:
            tuning=json.loads((ROOT/'Tools/ValidationData/Explosher/Tuning.json').read_text('utf-8'))
            values.update(tuning['weapons'])
            values.update(weaponTypeName='爆炸泼桶',moveSpeed=.088*60*18/24.037,swimSpeed=.1728*60*18/24.037)
    weapon_fields = re.findall(r'public (?:bool|float|int|string|double|WeaponFireMode|WeaponMuzzleMode|ProjectileMotionMode|AmmoConfigAsset) (\w+)\s*[;=]', (ROOT/'Assets/Splatoon/Config/WeaponConfigAsset.cs').read_text('utf-8-sig'))
    schema = ET.parse(ROOT/'Config/Luban/source/Defines/gameplay.xml')
    hero_fields = {node.attrib['name'] for bean in schema.iter('bean') if bean.attrib.get('name') == 'HeroConfig' for node in bean.findall('var')}
    check(hero_fields == set(columns), 'Source columns do not match Luban HeroConfig schema')
    check(not set(weapon_fields) & set(columns), 'Weapon fields remain in TbHero')
    check({r['id'] for r in rows} == {r['id'] for r in generated}, 'Source/generated hero ids differ')
    check({r['id'] for r in baseline} <= {r['id'] for r in rows}, 'Historical hero missing')
    check(len(columns) >= 31 and len(weapon_fields) >= 64, 'Unexpected split field counts')
    if args.hero is not None:
        rows = [r for r in rows if r['id'] == args.hero]
        check(len(rows) == 1, 'Requested hero missing or duplicated')
    group = (ROOT/'Assets/AddressableAssetsData/AssetGroups/Splatoon Local.asset').read_text('utf-8-sig')
    entries = dict(re.findall(r'  - m_GUID: (\w+)\s+ m_Address: (.+)', group))
    checked = 0
    for row in rows:
        gen = next(h for h in generated if h['id'] == row['id'])
        check(set(gen) == set(columns), f'Generated hero {row["id"]} contains extra/missing fields')
        for key, value in row.items(): check(equal(gen.get(key), value), f'Source/generated: {row["id"]}/{key}')
        path = ROOT / row['weaponConfigPath']
        check(path.exists(), f'Missing asset: {path}')
        if not path.exists(): continue
        values = asset_values(path)
        # Unity omits newly added fields until an asset is reserialized; runtime
        # defaults supply these optional modes. Explicit values must be known.
        check(set(values) <= set(weapon_fields), f'Unknown asset fields: {path}')
        meta = Path(str(path)+'.meta').read_text('utf-8')
        guid = re.search(r'^guid: (\w+)', meta, re.M)[1]
        check(entries.get(guid) == row['weaponConfigPath'], f'Addressables full path missing: {path}')
        check(list(entries.values()).count(row['weaponConfigPath']) == 1, f'Duplicate address: {path}')
        historical = next((h for h in baseline if h['id'] == row['id']), None)
        old = seconds_values(historical) if historical else {}
        if row['id'] == pistol_tuning['heroId']:
            old.update(pistol_tuning['preservedValues'])
            old.update(pistol_tuning['values'])
        if row['id'] == dualies_tuning['heroId']:
            old.update(dualies_tuning['values'])
        if row['id'] == blaster_tuning['heroId']:
            old.update(blaster_tuning['values'])
            old.update(blaster_tuning['heroValues'])
        aligned(row["id"], old)
        overrides = json.loads((ROOT/'Tools/ValidationData/PaintParity1130/VerifiedOverrides.json').read_text('utf-8'))
        old.update({r['field']: r['after'] for r in overrides['fields'] if r['hero'] == row['id']})
        combined = dict(gen, **values)
        for key, value in old.items():
            if key == 'displayName': continue
            checked += 1
            check(equal(combined.get(key), value), f'Migration changed original value: {row["id"]}/{key}')
        if historical and row['id'] not in (3, blaster_tuning['heroId']):
            check(values.get('spreadExpandSeconds') == 1 and values.get('spreadRecoverSeconds') == .5, f'New time defaults: {path}')
        if row['id'] == 7:
            bubble = dict(fireMode=5, motionMode=1, burstCount=4, pelletCount=1,
                          bubbleVolleySeconds=.55, bubbleIntervalSeconds=.05, shotInk=8,
                          startSeconds=.1, emergeStartSeconds=.2, inkRecoverLockSeconds=.65,
                          shootMoveSpeed=2.8, speedMin=14, speedMax=14, projectileGravity=18,
                          collisionRadius=.18, damage=30, damageMin=30, effectiveRange=24,
                          lifetime=2.4, bubbleGroundBounces=3, bubbleMaxBounces=6,
                          bubbleNormalRetention=.72, bubbleTangentRetention=.9, bubbleWallRetention=.9,
                          paintRadiusMin=.65, paintRadiusMax=.85, trailSpacing=.6,
                          trailRadiusMin=.2, trailRadiusMax=.3, trailMaxDrop=1,
                          spreadDegrees=0, jumpSpreadDegrees=0)
            aligned(7,bubble)
            for key, value in bubble.items():
                check(equal(combined.get(key), value), f'Bubble launch baseline: {key}')
        check(values.get('baseSpreadDegrees') == 0, f'Ground base: {path}')
        check(values.get('baseJumpSpreadDegrees') == 0, f'Air base: {path}')
    report = dict(passed=not errors, heroes=len(rows), characterFields=len(columns)-1, weaponPathFields=1,
                  originalWeaponFields=60, newWeaponFields=4, originalValuesCompared=checked, errors=errors)
    output = args.output if args.output is not None else ROOT/'Reports/WeaponAssets'
    output.mkdir(parents=True, exist_ok=True)
    name = f'hero-{args.hero}-static-validation.json' if args.hero is not None else 'static-validation.json'
    (output/name).write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n', 'utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if not errors else 1

if __name__ == '__main__': raise SystemExit(main())
