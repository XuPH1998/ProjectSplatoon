"""Apply the approved custom-hero balance to the authoritative Luban workbook."""
from pathlib import Path
from copy import copy
import json
import math
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
path = ROOT / 'Config/Luban/source/TbHero.xlsx'
wb = openpyxl.load_workbook(path)
ws = wb['Hero']
before = {(c.row, c.column): (c.value, copy(c._style)) for row in ws for c in row}
cols = {c.value: c.column for c in ws[1]}
additions = [('pelletCount','int','发射：每次有效发射的弹丸数量'), ('muzzleMode','int','发射：枪口模式（0单枪／1右左交替）'), ('semiBufferFrames','int','发射：半自动冷却末尾点击缓存（60Hz帧）')]
for key, kind, label in additions:
    if key not in cols:
        col = ws.max_column + 1
        cols[key] = col
        for row in range(1, 9):
            ws.cell(row,col)._style = copy(ws.cell(row,cols['burstCount'])._style)
            ws.cell(row,col).value = (key,kind,label)[row-1] if row <= 3 else (1 if key == 'pelletCount' else 0)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 32
common = dict(fireMode=3, burstCount=1, burstRecoveryFrames=0, semiBufferFrames=6)
changes = {
2:dict(common,name='DualPistolGirl',displayName='双持手枪',characterPrefabAddress='Character/DualPistolGirl',weaponPrefabAddress='Weapon/DualPistols',muzzleMode=1,fireIntervalFrames=10,fireRate=6,damage=32,damageMin=16,shotInk=.7,effectiveRange=8.4,speedMin=26,speedMax=26,spreadDegrees=3.5,jumpSpreadDegrees=8,damageReduceStartFrames=8,damageReduceEndFrames=40,shootMoveSpeed=4.2,inkRecoverLockFrames=15,paintRadiusMin=.55,paintRadiusMax=.7,trailRadius=.45),
3:dict(common,name='ShotgunGirl',displayName='霰弹枪',characterPrefabAddress='Character/ShotgunGirl',weaponPrefabAddress='Weapon/Shotgun',pelletCount=8,fireIntervalFrames=24,fireRate=2.5,damage=10,damageMin=4,shotInk=4,effectiveRange=6.5,speedMin=22,speedMax=22,spreadDegrees=7,jumpSpreadDegrees=11,damageReduceStartFrames=4,damageReduceEndFrames=18,shootMoveSpeed=3.2,inkRecoverLockFrames=30,paintRadiusMin=.24,paintRadiusMax=.32,trailRadius=.18),
4:dict(common,name='PistolGirl',displayName='精确手枪',characterPrefabAddress='Character/PistolGirl',weaponPrefabAddress='Weapon/Pistol',fireIntervalFrames=16,fireRate=3.75,damage=52,damageMin=26,shotInk=1.4,effectiveRange=12,speedMin=33,speedMax=33,spreadDegrees=1.5,jumpSpreadDegrees=6,damageReduceStartFrames=18,damageReduceEndFrames=42,shootMoveSpeed=3.4,inkRecoverLockFrames=22,paintRadiusMin=.4,paintRadiusMax=.55,trailRadius=.28),
5:dict(name='RocketLauncherGirl',displayName='蓄力火箭筒',characterPrefabAddress='Character/RocketLauncherGirl',weaponPrefabAddress='Weapon/RocketLauncher')}
measurement = ROOT / 'Docs/CombatGirls/FourHeroes/paint-measurements.json'
if measurement.exists():
    for item in json.loads(measurement.read_text('utf-8-sig')):
        if item['id'] in (2,3,4): changes[item['id']]['paintRange'] = round(item['range'], 2)
allowed = set()
for row in range(4,9):
    for key,value in changes.get(ws.cell(row,cols['id']).value,{}).items():
        allowed.add((row,cols[key])); ws.cell(row,cols[key]).value = value
for key,value in [('shotInk','发射：每次有效发射总耗墨（霰弹整组）'),('fireMode','发射：0全自动／1三连发／2蓄力／3半自动'),('paintRange','涂色：水平涂地参考距离（米，新半自动为实测值）')]:
    ws.cell(3,cols[key]).value=value; allowed.add((3,cols[key]))
for at,(value,style) in before.items():
    assert at in allowed or ws.cell(*at).value == value, at
    assert ws.cell(*at)._style == style, ('style',at)
expected = [[c.value for c in row] for row in ws]
wb.save(path); wb.close()
check = openpyxl.load_workbook(path)
for actual_row, expected_row in zip(check['Hero'], expected):
    for cell,value in zip(actual_row,expected_row):
        assert math.isclose(cell.value,value,rel_tol=1e-12,abs_tol=1e-12) if isinstance(value,(float,int)) else cell.value == value, cell.coordinate
assert check['Hero'].freeze_panes == 'E4'
check.close()
print(json.dumps(dict(heroes=5,fields=len(cols)-1,changedCells=len(allowed))))
