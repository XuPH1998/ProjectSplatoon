"""Install the pinned 11.3.0 Explosher profile; preserve workbook ZIP parts and GUIDs."""
from pathlib import Path
import argparse, hashlib, json, re, shutil, subprocess, urllib.request, zipfile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
DATA = ROOT / 'Tools/ValidationData/Explosher'
REPORT = ROOT / 'Reports/Explosher'
S = 18 / 24.037
URL = 'https://raw.githubusercontent.com/Leanny/splat3/7280ff9cde8bb1c5dcef46c700c326471584d2e6/data/parameter/1130/weapon/WeaponSlosherWashtub.game__GameParameterTable.json'

def prepare():
    DATA.mkdir(parents=True, exist_ok=True)
    REPORT.mkdir(parents=True, exist_ok=True)
    base = DATA / 'Baseline'
    if not base.exists():
        base.mkdir()
        paths = list((ROOT / 'Assets/GameResource/Weapons').glob('*/*Config.asset'))
        paths += [ROOT / 'Config/Luban/source/TbHero.xlsx', ROOT / 'Assets/GameResource/Bootstrap/Config/Luban/tbhero.json']
        manifest = {}
        for p in paths:
            rel = p.relative_to(ROOT).as_posix()
            manifest[rel] = hashlib.sha256(p.read_bytes()).hexdigest()
            if 'ShotgunGirl' in rel or p.name in ('TbHero.xlsx', 'tbhero.json'):
                shutil.copy2(p, base / p.name)
        (base / 'manifest.json').write_text(json.dumps(manifest, indent=2), encoding='utf8')
        (base / 'commit.txt').write_text(subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True), encoding='utf8')
    path = DATA / 'WeaponSlosherWashtub.1130.json'
    if not path.exists(): path.write_bytes(urllib.request.urlopen(URL).read())
    (DATA / 'Source.json').write_text(json.dumps({'url': URL, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'metresPerUnit': S}, indent=2), encoding='utf8')

def asset(path, updates):
    text = path.read_text(encoding='utf-8-sig')
    for key, value in updates.items():
        line = f'  {key}: {value}'
        if re.search(r'^  ' + key + ':', text, re.M): text = re.sub(r'^  ' + key + r':[^\n]*', lambda _: line, text, flags=re.M)
        else: text += line + '\n'
    path.write_text(text, encoding='utf8')

def install():
    updates = dict(fireMode=7, motionMode=5, fireRate=60/55, shotInk=11.7, startSeconds=16/60, emergeStartSeconds=22/60,
        inkRecoverLockSeconds=70/60, shootMoveSpeed=.045*60*S, pelletCount=1, semiBufferSeconds=.1,
        speedMin=1.4105*60*S, speedMax=1.4105*60*S, projectileGravity=.05*3600*S, lifetime=5,
        collisionRadius=.25*S, spreadDegrees=0, jumpSpreadDegrees=0, baseSpreadDegrees=0, baseJumpSpreadDegrees=0,
        straightSeconds=8/60, brakeSeconds=1/60, brakeSpeedMultiplier=1, effectiveRange=20.7*S,
        damage=55, damageMin=55, paintRadiusMin=2.28*S, paintRadiusMax=2.28*S,
        trailSpacing=3.1*S, trailRadiusMin=1.344*S, trailRadiusMax=1.344*S, referenceRules=1,
        referenceBrakeEndSpeed=10*60*S, referenceBrakeDrag=.1, referenceBrakeGravity=.04*3600*S,
        referenceFreeDrag=.12, referencePlayerRadius=.435*S, referenceTrailBudget=7,
        referenceTrailStart=3.1*S*.5, referenceTrailRandomPhase=0, referenceFootEvery=1, referenceFootRadius=1.56*S,
        paintDistanceMiddle=5*S, paintDistanceFar=15*S, paintDepthMin=1, paintDepthMax=1,
        paintDepthBreakMin=.7, paintDepthBreakMax=.7, paintBreakHeight=1.5*S, trailDepthScale=2,
        paintDropGravity=.016*3600*S, paintDropLifetime=5,
        wallDropRadius=1*S, wallDropGroundRadius=.7*S, wallDropSpeed=.08*60*S, wallDropSeconds=1.5,
        collisionExplosionPaintRadius=4.5*S, explosherAirSpeed=1.38975*60*S,
        explosherUpwardRate=.1, explosherMoveSideRate=.3, explosherMoveForwardRate=1, explosherMoveVerticalRate=.5,
        explosherFieldInitialRadius=.1*S, explosherPlayerInitialRadius=.0435*S,
        explosherFieldGrowSeconds=4/60, explosherPlayerGrowSeconds=5/60,
        explosherPostSeconds=35/60, explosherMoveLimitSeconds=55/60,
        explosherPaintNearDistance=21*S, explosherPaintFarDistance=23*S, explosherBlastOffset=.2*S,
        explosherTrailPhaseMax=.55, explosherFootDepth=1.4)
    asset(ROOT / 'Assets/GameResource/Weapons/ShotgunGirl/ShotgunGirlWeaponConfig.asset', updates)
    asset(ROOT / 'Assets/GameResource/Weapons/ShotgunGirl/ShotgunGirlAmmoConfig.asset', dict(
        explosionEnabled=1, explosionRadius=2.87*S, explosionDamage=35, explosionConstantDamage=1,
        collisionExplosionRadiusRate=1, collisionExplosionDamageRate=1, excludeDirectHitFromExplosion=0,
        explosionPaint=1, explosionPaintRadiusMin=4.095*S, explosionPaintRadiusMax=4.5*S,
        visualLifetime=5, particlesPerBurst=1, satelliteSpread=0, maxRibbonGap=1.1))
    # Edit only three cell elements; every unrelated ZIP member stays byte-identical.
    path = ROOT / 'Config/Luban/source/TbHero.xlsx'
    import openpyxl
    wb = openpyxl.load_workbook(path)
    sheet = wb.active
    cols = {c.value: c.column for c in sheet[1]}
    row = next(r for r in range(5, sheet.max_row+1) if sheet.cell(r, cols['id']).value == 3)
    changes = {'weaponTypeName': '爆炸泼桶', 'moveSpeed': .088*60*S, 'swimSpeed': .1728*60*S}
    cells = {sheet.cell(row, cols[k]).coordinate: v for k,v in changes.items()}
    with zipfile.ZipFile(path) as z: parts = [(i, z.read(i.filename)) for i in z.infolist()]
    for i,(info, data) in enumerate(parts):
        if info.filename != 'xl/worksheets/sheet1.xml': continue
        text = data.decode('utf8')
        for cell, value in cells.items():
            pattern = r'<c\b(?=[^>]*\br="'+cell+r'")[^>]*>.*?</c>'
            old = re.search(pattern, text).group(0)
            style = re.search(r'\bs="[^"]*"', old)
            attrs = ' ' + style.group(0) if style else ''
            payload = f't="inlineStr"><is><t>{value}</t></is>' if isinstance(value,str) else f't="n"><v>{value}</v>'
            text = text.replace(old, f'<c r="{cell}"{attrs} {payload}</c>')
        parts[i] = (info, text.encode('utf8'))
    with zipfile.ZipFile(path,'w') as z:
        for info,data in parts: z.writestr(info,data)
    tuning = json.dumps({'weapons':updates,'heroCells':cells,'adaptations':[
        'Space scale follows existing project weapons.', 'Emergence adds 6 frames; dive lock is 35 frames.',
        'Brake phase is 1 reference frame; inherited default semantics are project reference integration.',
        'Wall drips use existing constant-speed surface simulation; no global knockback changes.',
        'Paint shape, drop gravity/lifetime and safety lifetime use project renderer/physics adaptations.'
    ]}, ensure_ascii=False, indent=2)
    (REPORT / 'tuning.json').write_text(tuning, encoding='utf8')
    (DATA / 'Tuning.json').write_text(tuning, encoding='utf8')
    print('Installed Explosher assets and hero cells:', cells)

if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--apply', action='store_true'); args = parser.parse_args()
    prepare()
    if args.apply: install()
    else: print('Baseline and pinned source ready')
