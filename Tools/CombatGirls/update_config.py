"""Change only the default character/weapon identity cells; preserve gameplay values."""
from pathlib import Path
import argparse
import json
import re
import zipfile
from xml.sax.saxutils import escape
import openpyxl

parser = argparse.ArgumentParser()
parser.add_argument("--project", default=str(Path(__file__).resolve().parents[2]))
args = parser.parse_args()
for file, name, address, address_key in [
    ("TbCharacter.xlsx", "RifleGirl", "Character/RifleGirl", "visualAddress"),
    ("TbWeapon.xlsx", "RifleGirlRifle", "Weapon/RifleGirlRifle", "prefabAddress"),
]:
    path = Path(args.project) / "Config/Luban/source" / file
    wb = openpyxl.load_workbook(path)
    sheet = wb.worksheets[0]
    before = {cell.coordinate: cell.value for row in sheet for cell in row}
    header = next(row for row in sheet if "id" in [c.value for c in row] and "name" in [c.value for c in row])
    cols = {c.value: c.column for c in header if c.value is not None}
    row = next(r for r in range(header[0].row + 1, sheet.max_row + 1) if sheet.cell(r, cols["id"]).value == 1)
    changed = []
    for key, value in (("name", name), (address_key, address)):
        cell = sheet.cell(row, cols[key])
        changed.append({"cell": cell.coordinate, "before": cell.value, "after": value})
        cell.value = value
    expected = {c.coordinate: c.value for r in sheet for c in r}
    allowed = {x["cell"] for x in changed}
    assert all(expected[c] == v or c in allowed for c, v in before.items())
    temp = path.with_suffix(".combatgirls.tmp.xlsx")
    # Preserve every unrelated XML entry, style and empty-string cell exactly.
    with zipfile.ZipFile(path) as original, zipfile.ZipFile(temp, 'w') as output:
        for entry in original.infolist():
            data = original.read(entry.filename)
            if entry.filename == 'xl/worksheets/sheet1.xml':
                xml = data.decode('utf-8')
                for change in changed:
                    pattern = r'<c\b([^>]*\br="' + change['cell'] + r'"[^>]*)>(.*?)</c>'
                    match = re.search(pattern, xml, re.S)
                    assert match, (file, change['cell'])
                    attrs = re.sub(r'\s+t="[^"]*"', '', match[1])
                    replacement = '<c' + attrs + ' t="inlineStr"><is><t>' + escape(change['after']) + '</t></is></c>'
                    xml = xml[:match.start()] + replacement + xml[match.end():]
                data = xml.encode('utf-8')
            output.writestr(entry, data)
    check = openpyxl.load_workbook(temp)
    actual = {c.coordinate: c.value for r in check.worksheets[0] for c in r}
    assert actual == expected, (file, "round-trip changed cell contents")
    check.close()
    temp.replace(path)
    print(json.dumps({"file": file, "changes": changed}, ensure_ascii=False))
