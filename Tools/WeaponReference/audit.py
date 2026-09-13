"""Rebuild the complete reference ledger without changing any runtime data.

Use the bundled Python (openpyxl is used only to read the authoritative workbooks).
--check verifies source/generated agreement and the approved scalar change set.
"""
from pathlib import Path
import argparse, csv, hashlib, json, math, subprocess
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
DOC = ROOT / 'Docs/WeaponAudit'
BASE = '00be31ff3662c77cb2b62d1f6f61ce3c62492d19'
LEAN = '7280ff9cde8bb1c5dcef46c700c326471584d2e6'
SENDOU = '55889eccc7c18f24569098959ae28f4a7034ce4e'
NAMES = {1:'WeaponShooterNormal',2:'WeaponShooterBlaze',3:'WeaponShooterGravity',4:'WeaponShooterTripleQuick',5:'WeaponChargerNormal'}
CHANGES = {('Weapon',2,'shotInk'):.5, ('Weapon',2,'inkRecoverLockFrames'):15,
           ('Weapon',3,'shotInk'):1.5, ('Weapon',3,'fireIntervalFrames'):9, ('Weapon',3,'fireRate'):60/9,
           ('Weapon',4,'shotInk'):1.1, ('Weapon',4,'inkRecoverLockFrames'):25,
           ('Weapon',5,'damage'):160, ('Character',1,'swimRecoverInk'):100*60/180}
FIELDS = {
    'fireRate':'WeaponParam.RepeatFrame','fireIntervalFrames':'WeaponParam.RepeatFrame',
    'damage':'DamageParam.ValueMax','damageMin':'DamageParam.ValueMin', 'shotInk':'WeaponParam.InkConsume',
    'speedMin':'MoveParam.SpawnSpeed','speedMax':'MoveParam.SpawnSpeed','gravity':'MoveParam.FreeGravity',
    'spreadDegrees':'WeaponParam.Stand_DegSwerve','jumpSpreadDegrees':'WeaponParam.Jump_DegSwerve',
    'inkRecoverLockFrames':'WeaponParam.InkRecoverStop','damageReduceStartFrames':'DamageParam.ReduceStartFrame',
    'damageReduceEndFrames':'DamageParam.ReduceEndFrame','straightFrames':'MoveParam.GoStraightToBrakeStateFrame',
    'shootMoveSpeed':'WeaponParam.MoveSpeed','burstRecoveryFrames':'WeaponParam.TripleShotSpanFrame',
    'chargeFrames':'WeaponParam.ChargeFrameFullCharge','chargeMinDamage':'DamageParam.ValueMinCharge',
    'chargePartialMaxDamage':'DamageParam.ValueMaxCharge','chargeMinInk':'WeaponParam.InkConsumeMinCharge',
    'chargeMinRange':'MoveParam.DistanceMinCharge','chargeMinSpeed':'MoveParam.SpawnSpeedMinCharge',
}
GROUPS = {
    'metadata': ('id','name','displayName','prefabAddress'),
    'paint': ('paintRadiusMin','paintRadiusMax','paintHardness','paintStrength','trailSpacing','trailRadius','trailMaxDrop','paintRange','chargeMinPaintRange'),
    'ballistics': ('speedMin','speedMax','gravity','lifetime','collisionRadius','straightFrames','brakeFrames','brakeSpeedMultiplier','effectiveRange','chargeMinRange','chargeMinSpeed'),
    'damage': ('damage','damageMin','damageReduceStartFrames','damageReduceEndFrames','chargeMinDamage','chargePartialMaxDamage'),
    'dispersion': ('spreadDegrees','jumpSpreadDegrees','spreadRecoverFrames','chargeMinSpread','chargeMinJumpSpread'),
    'ink': ('shotInk','inkRecoverLockFrames','chargeMinInk'),
    'movement': ('shootMoveSpeed',),
}
REASONS = {
    'metadata':'项目标识/资源，不属于原作玩法数值',
    'paint':'缺少原作轮廓/落墨生成及下落算法、统一空间换算；硬度/强度无直接对应',
    'ballistics':'缺少可核实的统一空间标定和完整状态转移/碰撞规则；保留整段旧模型',
    'damage':'原始端点可读，但未取得固定版本衰减/蓄力插值及取整算法；保留整段旧曲线',
    'dispersion':'缺少角度口径、随机分布、偏差累积与恢复算法；保留整段旧模型',
    'ink':'缺失默认值或蓄力耗墨/扣墨时机规则未闭合',
    'movement':'空间比例、重量默认分类和状态限制未闭合',
    'firing':'起手/组间/蓄力状态规则及继承默认值未闭合',
}
CONSUMERS = {
    'metadata':'GameplayConfig.Validate / PrototypePlayer.Equipment / WeaponDisplay',
    'paint':'InkProjectileService.Resolve/PaintTrail; InkBrush.Coverage; PaintSurface; InkTexturePainter.shader',
    'ballistics':'InkBallistics.Position/TravelTime/LaunchVelocity; InkProjectileService.Simulate/Resolve; InkPresentation.LateUpdate',
    'damage':'WeaponSimulation.Damage -> InkProjectileService.Resolve -> PrototypePlayer.ReceiveDamage',
    'dispersion':'WeaponSimulation.Spread; PrototypePlayer.Network.Step; InkBallistics.LaunchVelocity',
    'ink':'WeaponSimulation.InkCost/Emit; ResourceSimulation.Step',
    'movement':'PrototypePlayer.Network.Step -> PlayerMotorSimulation.Step',
    'firing':'WeaponSimulation.Step/Begin/Emit; WeaponDisplay.Cadence',
}
def load(path): return json.loads(path.read_text(encoding='utf-8-sig'))
def base(table):
    path=f'Assets/GameResource/Bootstrap/Config/Luban/tb{table.lower()}.json'
    return json.loads(subprocess.check_output(['git','show',f'{BASE}:{path}'],cwd=ROOT))
def hero_field(table, field):
    mapping=load(ROOT/'Docs/HeroMigration/Field-Mapping.json')
    return next((m['heroField'] for m in mapping if m['sourceTable']==table and m['sourceField']==field),None)

def project_legacy(table, rows):
    # Historical reference ledgers keep their old field names; all live data comes from TbHero.
    if table in ('Character','Weapon'):
        before=base(table); projected=[]
        for i,row in enumerate(rows if table=='Weapon' else rows[:1]):
            template=before[i] if table=='Weapon' else before[0]
            projected.append({field:row[hero_field(table,field)] if hero_field(table,field) else value for field,value in template.items()})
        return projected
    if table=='RoomMode':
        previous={r['id']:r for r in base(table)}
        return [{**{k:v for k,v in row.items() if k!='heroId'},'weaponId':row['heroId'],'characterId':previous[row['id']]['characterId']} for row in rows]
    return rows

def current(table):
    actual='Hero' if table in ('Character','Weapon') else table
    return project_legacy(table,load(ROOT/f'Assets/GameResource/Bootstrap/Config/Luban/tb{actual.lower()}.json'))
def equal(a,b): return math.isclose(a,b,rel_tol=1e-6,abs_tol=1e-6) if isinstance(a,(float,int)) and isinstance(b,(float,int)) else a==b
def get(node,path):
    for key in path.split('.'):
        if not isinstance(node,dict) or key not in node: return None
        node=node[key]
    return node
def flatten(node,prefix=''):
    for key,value in node.items():
        path=prefix+key
        if isinstance(value,dict): yield from flatten(value,path+'.')
        else: yield path,value
def text(value): return '待核实（覆盖中缺失，不等于零）' if value is None else json.dumps(value,ensure_ascii=False)
def source(table):
    if table in ('Character','Weapon'):
        rows,labels=source('Hero')
        return project_legacy(table,rows),{k:labels.get(hero_field(table,k),'历史角色标识，已并入英雄身份') for k in base(table)[0]}
    wb=openpyxl.load_workbook(ROOT/f'Config/Luban/source/Tb{table}.xlsx',data_only=False)
    sheet=wb[table]; headers=[c.value for c in sheet[1]][1:]
    rows=[dict(zip(headers,[c.value for c in row][1:])) for row in sheet.iter_rows(min_row=4) if row[1].value is not None]
    labels={key:sheet.cell(3,i+2).value for i,key in enumerate(headers)}
    wb.close()
    if table=='RoomMode':
        rows=project_legacy(table,rows);labels['weaponId']=labels.pop('heroId');labels['characterId']='旧版角色引用（合表前）'
    return rows,labels
def verify():
    changes=[]
    for table in ['Weapon','Character','Global','Map','RoomMode']:
        rows,_=source(table); generated=current(table); before={row['id']:row for row in base(table)}
        assert len(rows)==len(generated),(table,'row count')
        for row,saved in zip(rows,generated):
            assert set(row)==set(saved),(table,'schema')
            for key,value in row.items():
                assert equal(value,saved[key]),(table,row['id'],key,'source/generated mismatch')
                signature=(table,row['id'],key)
                if signature in CHANGES:
                    assert equal(value,CHANGES[signature]),(signature,'verified scalar changed')
                elif not equal(value,before[row['id']][key]): raise AssertionError((signature,'unapproved change to gated baseline'))
                if not equal(value,before[row['id']][key]): changes.append({'table':table,'id':row['id'],'field':key,'before':before[row['id']][key],'after':value})
    assert len(changes)==len(CHANGES)
    return changes
def write_csv(path,columns,rows):
    with path.open('w',encoding='utf-8-sig',newline='') as stream:
        writer=csv.DictWriter(stream,fieldnames=columns); writer.writeheader(); writer.writerows(rows)
def verify_snapshots():
    manifest=load(DOC/'Implementation-Manifest.json')
    assert manifest['referenceVersion']=='11.3.0' and manifest['leanCommit']==LEAN
    expected={f'{name}.1130.json' for name in NAMES.values()} | {'Common.1130.json'}
    assert set(manifest['referenceSha256'])==expected
    for name,digest in manifest['referenceSha256'].items():
        assert hashlib.sha256((DOC/name).read_bytes()).hexdigest()==digest,(name,'pinned reference hash mismatch')

def generate():
    if (DOC/'Implementation-Manifest.json').exists(): verify_snapshots()
    changes=verify(); ledger=[]; originals=[]
    _,labels=source('Weapon')
    before={row['id']:row for row in base('Weapon')}
    for weapon in current('Weapon'):
        id=weapon['id']; name=NAMES[id]; raw=load(DOC/f'{name}.1130.json')['GameParameters']
        refs={}
        for field,value in weapon.items():
            group=next((g for g,keys in GROUPS.items() if field in keys),'firing')
            path=FIELDS.get(field)
            if id==5:
                path={'damage':'DamageParam.ValueFullCharge','damageMin':None,'shotInk':'WeaponParam.InkConsumeFullCharge',
                      'speedMin':'MoveParam.SpawnSpeedFullCharge','speedMax':'MoveParam.SpawnSpeedFullCharge',
                      'shootMoveSpeed':'WeaponParam.MoveSpeedFullCharge','effectiveRange':'MoveParam.DistanceFullCharge'}.get(field,path)
            raw_value=get(raw,path) if path else None
            if path: refs.setdefault(path,[]).append(hero_field('Weapon',field) or field)
            status='保留旧机制'; reason=REASONS[group]; conversion='未确认或无一一对应关系'
            if group=='metadata': status='项目配置'; conversion='不适用'
            if field in ('damage','damageMin','chargeMinDamage','chargePartialMaxDamage'): conversion='原作伤害 / 10 = 项目 HP；曲线与取整另行核实'
            if field in ('shotInk','chargeMinInk'): conversion='原作整罐比例 × 100 = 项目墨点；蓄力曲线另行核实'
            if field in ('fireIntervalFrames','inkRecoverLockFrames','damageReduceStartFrames','damageReduceEndFrames','straightFrames','chargeFrames'): conversion='原作参考帧 / 60 = 秒；状态边界另行核实'
            if field=='fireRate': conversion='60 / RepeatFrame；仅配置校验；HUD 射速由实际帧间隔推导'
            if ('Weapon',id,field) in CHANGES:
                status='已修改：独立标量'; reason='仅启用已确认的单次耗墨、间隔、锁定或满蓄离散伤害，不声明所在完整机制已复刻'
            elif id<5 and field=='shotInk' and raw_value is not None and equal(value,raw_value*100):
                status='数值已一致：单次耗墨'; reason='每发固定扣墨，100 点墨量映射已确认'
            elif id==2 and field=='fireIntervalFrames': status='数值已一致：持续间隔'; reason='显式 RepeatFrame=4'
            if field in ('paintRange','chargeMinPaintRange'): reason='只作为水平涂地标定目标；未参与运行时截断；无合格原作测量样本'
            if id==5 and field=='damageMin': reason='当前蓄力枪不消费该字段；保留旧值，不能当作满蓄伤害'
            consumer=CONSUMERS[group]
            if field in ('paintRange','chargeMinPaintRange'): consumer='GameplayConfig.Validate; WeaponDisplay.Details; 编辑器测量目标（不参与实际涂色）'
            if field=='fireRate': consumer='GameplayConfig.Validate（不直接参与发射和 HUD 射速）'
            if id==5 and field=='damageMin': consumer='GameplayConfig.Validate；蓄力伤害分支不消费'
            ledger.append(dict(weapon=id,name=weapon['displayName'],group=group,field=field,heroField=hero_field('Weapon',field),meaning=labels[field],baseline=text(before[id][field]),current=text(value),referencePath='GameParameters.'+path if path else '无直接字段',referenceValue=text(raw_value),conversion=conversion,consumer=consumer,status=status,reason=reason,source=f'https://raw.githubusercontent.com/Leanny/splat3/{LEAN}/data/parameter/1130/weapon/{name}.game__GameParameterTable.json'))
        for path,value in flatten(raw):
            if path.endswith('$type'): continue
            originals.append(dict(weapon=id,path='GameParameters.'+path,value=text(value),projectFields=','.join(refs.get(path,[])) or '未建模/无直接对应',status='原始覆盖；不包含缺失的类型默认值'))
    write_csv(DOC/'Full-Parameter-Coverage.csv',list(ledger[0]),ledger)
    write_csv(DOC/'Raw-Parameter-Coverage.csv',list(originals[0]),originals)
    common=[]; common_raw=load(DOC/'Common.1130.json')
    common_refs={'recoverInk':'InkRecoverFrm_Std','swimRecoverInk':'InkRecoverFrm_Stealth',
                 'moveSpeed':'MoveVel_Human','swimSpeed':'MoveVel_Stealth'}
    for table in ['Character','Global']:
        _,labels=source(table); before={row['id']:row for row in base(table)}
        for row in current(table):
            for field,value in row.items():
                reference=common_refs.get(field) if table=='Character' else None
                raw_value=common_raw[reference][2] if reference else None
                status='保留项目基线；关联机制未通过证据门槛'
                conversion='无已确认的对应换算'
                reason='未取得固定版本算法、默认值和完整单位依据；不能把缺失当成零'
                consumer='PlayerMotorSimulation.Step / GameplayConfig.Validate' if table=='Character' else 'GameplayConfig.Validate / PrototypeApp / PrototypeMatch'
                if field in ('recoverInk','swimRecoverInk'):
                    status='已修改：180 帧潜墨回满' if field=='swimRecoverInk' else '数值已一致：600 帧人形回满'
                    conversion='100 × 60 / 原作回满帧数 = 点/秒；零能力点取 [High,Mid,Low] 的 Low'
                    reason='只确认允许回墨后的固定恢复速率，不证明所有回墨状态转移'
                    consumer='ResourceSimulation.Step -> PrototypeRules.Recover'
                elif reference:
                    reason='显式值仅为标准重量零能力点线索；Fast/Slow 分支、重量归类和统一空间映射未闭合'
                elif field in ('id','name','visualAddress') or (table=='Global' and not field.startswith('paint')):
                    status='项目配置'; reason='项目资源或工程参数，无需按原作逐值替换'
                if table=='Global' and field.startswith('paint'):
                    consumer='InkBrush.Coverage / SurfaceOwnershipGrid / InkTexturePainter.shader'
                    reason='项目覆盖算法参数，未取得原作一一对应依据'
                if table=='Character' and ('Health' in field or field in ('healthRecoverDelay','healthRecoverRate','maxHealth','enemyInkDamageRate')):
                    consumer='ResourceSimulation.Step / PrototypePlayer.ReceiveDamage / GameplayConfig.Validate'
                if field=='maxInk':
                    consumer='ResourceSimulation.Step / WeaponSelectionRules / PrototypePlayer.Network / HUD'
                    reason='项目 100 点表示整罐；原作整罐比例 × 100 的换算已确认'
                common.append(dict(table=table,field=field,heroField=hero_field(table,field) if table=='Character' else '',meaning=labels[field],baseline=text(before[row['id']][field]),current=text(value),
                    referencePath=reference+'[2]' if reference else '未确认直接字段',referenceValue=text(raw_value),
                    referenceKind='公开能力曲线的零能力点值' if reference else '无已确认对应；不是零',conversion=conversion,
                    consumer=consumer,status=status,reason=reason,
                    source=f'https://raw.githubusercontent.com/Leanny/splat3/{LEAN}/data/parameter/1130/misc/params.json' if reference else ''))
    write_csv(DOC/'Common-Parameter-Coverage.csv',list(common[0]),common)
    hashes={p.name:hashlib.sha256(p.read_bytes()).hexdigest() for p in [*(DOC/f'{name}.1130.json' for name in NAMES.values()),DOC/'Common.1130.json']}
    manifest={'referenceVersion':'11.3.0','baselineCommit':BASE,'leanCommit':LEAN,'sendouReferenceCommit':SENDOU,'completeReplica':False,'runtimeTable':'TbHero','ledgerFieldNames':'Historical projection; heroField names the live merged field.',
              'changes':changes,'referenceSha256':hashes,'gatedMechanisms':REASONS,
              'sourceCounts':{'projectWeaponFields':len(ledger),'rawReferenceFields':len(originals),'commonFields':len(common)}}
    (DOC/'Implementation-Manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(f'Ledger: {len(ledger)} weapon fields, {len(originals)} raw fields, {len(common)} common fields; {len(changes)} verified scalar cells.')
if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--check',action='store_true'); args=parser.parse_args()
    if args.check:
        verify_snapshots()
        print(f'PASS: source/generated data agree; {len(verify())} approved cells; gated values unchanged; 6 pinned reference hashes match.')
    else: generate()
