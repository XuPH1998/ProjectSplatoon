"""Verify the committed pre-merge numeric baseline against authoritative Hero XLSX and generated JSON."""
from pathlib import Path
import argparse,csv,json,math,hashlib
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
def load(path):return json.loads(path.read_text(encoding='utf-8-sig'))
def equal(a,b):return math.isclose(a,b,rel_tol=1e-6,abs_tol=1e-6) if isinstance(a,(int,float)) and isinstance(b,(int,float)) else a==b
def verify():
    docs=ROOT/'Docs/HeroMigration';baseline=load(docs/'Migration-Baseline.json');mapping=load(docs/'Field-Mapping.json')
    path=ROOT/'Assets/GameResource/Bootstrap/Config/Luban'
    heroes=load(path/'tbhero.json');weapons={r['id']:r for r in baseline['Weapon']}
    assert len(heroes)==5 and {h['id'] for h in heroes}==set(weapons)
    for h in heroes:
        assert set(h)=={m['heroField'] for m in mapping}
        for m in mapping:
            previous=weapons[h['id']] if m['sourceTable']=='Weapon' else baseline['Character'][0]
            assert equal(h[m['heroField']],previous[m['sourceField']]),(h['id'],m['heroField'],'baseline changed')
    mode=load(path/'tbroommode.json');expected=[{('heroId' if k=='weaponId' else k):v for k,v in r.items() if k!='characterId'} for r in baseline['RoomMode']]
    assert mode==expected
    for table in ['Hero','RoomMode','Map','Global']:
        wb=openpyxl.load_workbook(ROOT/f'Config/Luban/source/Tb{table}.xlsx');ws=wb[table]
        headers=[c.value for c in ws[1]][1:];data=[dict(zip(headers,[c.value for c in row][1:])) for row in ws.iter_rows(min_row=4) if row[1].value is not None]
        saved=load(path/f'tb{table.lower()}.json');assert len(data)==len(saved)
        for a,b in zip(data,saved):
            assert set(a)==set(b)
            for k in a:assert equal(a[k],b[k]),(table,k,'source/generated mismatch')
        wb.close()
    for old in ['Character','Weapon']:
        assert not (ROOT/f'Config/Luban/source/Tb{old}.xlsx').exists()
        assert not (path/f'tb{old.lower()}.json').exists()
        assert not (ROOT/f'Assets/Splatoon/Config/Generated/{old}Config.cs').exists()
    group=(ROOT/'Assets/AddressableAssetsData/AssetGroups/Splatoon Local.asset').read_text('utf-8-sig')
    assert 'm_Address: tbhero\n' in group and 'm_Address: tbweapon\n' not in group and 'm_Address: tbcharacter\n' not in group
    return {'heroes':len(heroes),'fieldsPerHero':len(mapping),'verifiedValues':len(heroes)*len(mapping),'baselineCommit':baseline['commit']}
def verify_measurements(directory):
    records=list(csv.DictReader((ROOT/'Docs/HeroMigration/Measurement-Comparison.csv').open(encoding='utf-8-sig')))
    assert len(records)==87
    for record in records:
        name=record['scenario']
        assert load(directory/(name+'.json'))['gridHash']==record['gridHash'],name
        for suffix in ['trace.csv','paint.csv']:
            assert hashlib.sha256((directory/(name+'.'+suffix)).read_bytes()).hexdigest()==record[suffix+' SHA256'],name+'.'+suffix
    return len(records)

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--measurements',type=Path,help='Compare captured measurements with pre-migration trajectory, paint and ownership hashes.')
    args=parser.parse_args();result=verify()
    if args.measurements:result['unchangedMeasurementScenarios']=verify_measurements(args.measurements)
    print('PASS:',json.dumps(result,ensure_ascii=False))
