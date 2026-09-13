"""Extend the existing Luban source workbook; preserve existing sheets and styles."""
from pathlib import Path
from copy import copy
import re
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
FIELDS = [
    ('displayName', 'string', '枪械中文显示名', '标准射击枪'),
    ('fireMode', 'int', '机制：0全自动／1三连发／2蓄力松开发射', 0),
    ('shootMoveSpeed', 'float', '射击或蓄力移动速度（米/秒）', 3.6),
    ('burstCount', 'int', '每组发数（全自动为1）', 1),
    ('burstRecoveryFrames', 'int', '末发到下一组首发（60Hz参考帧）', 0),
    ('chargeFrames', 'int', '满蓄时间（60Hz参考帧，非蓄力为0）', 0),
    ('chargeMinDamage', 'float', '点射伤害（HP）', 0),
    ('chargePartialMaxDamage', 'float', '未满蓄伤害上限（HP）', 0),
    ('chargeMinInk', 'float', '点射耗墨（点）', 0),
    ('chargeMinRange', 'float', '点射伤害射程（米）', 0),
    ('chargeMinSpeed', 'float', '点射初速（米/秒）', 0),
    ('chargeMinSpread', 'float', '点射地面散布半角（度）', 0),
    ('chargeMinJumpSpread', 'float', '点射空中散布半角（度）', 0),
    ('chargeMinPaintRange', 'float', '点射水平涂地射程目标（米）', 0),
]
path = ROOT/'Config/Luban/source/TbWeapon.xlsx'
wb = openpyxl.load_workbook(path)
ws = wb.worksheets[0]
cols = {c.value: c.column for c in ws[1] if c.value}
for key, kind, comment, default in FIELDS:
    if key not in cols:
        col = ws.max_column + 1
        cols[key] = col
        for row, value in enumerate([key, kind, comment, default], 1):
            cell = ws.cell(row, col, value)
            cell._style = copy(ws.cell(row, col - 1)._style)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 28
    ws.cell(4, cols[key], default)
# Explicitly authorized restoration of the user's later gravity experiment.
ws.cell(4, cols['gravity'], 9.8)
base = {key: ws.cell(4, col).value for key, col in cols.items()}
overrides = [
    dict(id=2, name='LightShooter', displayName='轻型速射枪', fireIntervalFrames=4,
         fireRate=15, damage=24, damageMin=12, shotInk=.55, effectiveRange=8.4,
         shootMoveSpeed=4, spreadDegrees=7, jumpSpreadDegrees=14, speedMin=26, speedMax=26, paintRange=11.6),
    dict(id=3, name='HeavyShooter', displayName='重型射击枪', fireIntervalFrames=10,
         fireRate=6, damage=52, damageMin=26, shotInk=1.8, effectiveRange=12,
         shootMoveSpeed=3, spreadDegrees=4, jumpSpreadDegrees=10, speedMin=34, speedMax=34, paintRange=15.2),
    dict(id=4, name='BurstShooter', displayName='三连发射击枪', fireMode=1, burstCount=3,
         burstRecoveryFrames=16, fireIntervalFrames=4, fireRate=15, damage=34, damageMin=17,
         shotInk=.9, effectiveRange=12, shootMoveSpeed=3.4, spreadDegrees=3,
         jumpSpreadDegrees=9, speedMin=33, speedMax=33, paintRange=14.8),
    dict(id=5, name='ChargeShooter', displayName='蓄力狙击枪', fireMode=2, chargeFrames=60,
         fireIntervalFrames=30, fireRate=2, damage=120, damageMin=120, shotInk=8,
         effectiveRange=18, shootMoveSpeed=2.4, spreadDegrees=.2, jumpSpreadDegrees=1,
         speedMin=46, speedMax=46, paintRange=20.4, inkRecoverLockFrames=45,
         chargeMinDamage=40, chargePartialMaxDamage=80, chargeMinInk=2, chargeMinRange=8,
         chargeMinSpeed=25, chargeMinSpread=2, chargeMinJumpSpread=6, chargeMinPaintRange=11.2),
]
for values in overrides:
    row = next((r for r in range(4, ws.max_row + 1) if ws.cell(r, cols['id']).value == values['id']), ws.max_row + 1)
    for key, value in (base | values).items():
        cell = ws.cell(row, cols[key], value)
        cell._style = copy(ws.cell(4, cols[key])._style)
    ws.row_dimensions[row].height = ws.row_dimensions[4].height
wb.save(path)
schema_path = ROOT/'Config/Luban/source/Defines/gameplay.xml'
schema = schema_path.read_text(encoding='utf-8')
match = re.search(r'(<bean name="WeaponConfig"[^>]*>)(.*?)(  </bean>)', schema, re.S)
body = match[2]
for key, kind, comment, default in FIELDS:
    if f'name="{key}"' not in body:
        body += f'    <var name="{key}" type="{kind}" comment="{comment}"/>\n'
schema_path.write_text(schema[:match.start()]+match[1]+body+match[3]+schema[match.end():], encoding='utf-8')
check = openpyxl.load_workbook(path, read_only=True, data_only=True)
assert [check.worksheets[0].cell(r, cols['id']).value for r in range(4,9)] == [1,2,3,4,5]
print('Five weapon source rows and schema updated; existing workbook layout preserved.')
