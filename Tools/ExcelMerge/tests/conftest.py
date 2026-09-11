from __future__ import annotations

import tempfile
from pathlib import Path

import openpyxl
import pytest


@pytest.fixture
def tmp_dir():
    with tempfile.TemporaryDirectory() as d:
        yield Path(d)


def create_xlsx(path: Path, sheets: dict[str, list[list]]) -> Path:
    """Create an xlsx file from a dict of sheet_name -> 2D list of values."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(title=name)
        for r_idx, row in enumerate(rows, 1):
            for c_idx, val in enumerate(row, 1):
                if val is not None:
                    ws.cell(row=r_idx, column=c_idx, value=val)
    wb.save(str(path))
    wb.close()
    return path
