import openpyxl
wb = openpyxl.load_workbook("demo_styled/merged.xlsx")
ws = wb["ItemEnum"]

c = ws.cell(row=1, column=1)
print(f"A1 value: {c.value}")
print(f"A1 fill: {c.fill.start_color.rgb}")
print(f"A1 font bold: {c.font.bold}, color: {c.font.color.rgb}")
print(f"A1 border: {c.border.left.style}")

c2 = ws.cell(row=2, column=8)
print(f"H2 value: {c2.value}")
print(f"H2 fill: {c2.fill.start_color.rgb}")

print(f"J6 (conflict->ours): {ws.cell(row=6, column=10).value}")
print(f"I8 (theirs alias): {ws.cell(row=8, column=9).value}")
print(f"H9 (ours new row): {ws.cell(row=9, column=8).value}")

col_a = ws.column_dimensions["A"].width
col_h = ws.column_dimensions["H"].width
print(f"Col A width: {col_a}")
print(f"Col H width: {col_h}")
