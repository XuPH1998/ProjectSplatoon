import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side

os.makedirs("demo_styled", exist_ok=True)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "ItemEnum"

header_fill = PatternFill(start_color="8B4513", end_color="8B4513", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True)
green_fill = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid")
thin_border = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

headers = ["##var", "full_name", "flags", "unique", "group", "comment", "tags"]
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.border = thin_border

data_headers = ["name", "alias", "value", "comment"]
for c, h in enumerate(data_headers, 8):
    cell = ws.cell(row=2, column=c, value=h)
    cell.fill = green_fill
    cell.border = thin_border

ws.cell(row=4, column=2, value="ItemEnum")
ws.cell(row=4, column=4, value=False)
ws.cell(row=4, column=8, value="None")
ws.cell(row=4, column=9, value="无")
ws.cell(row=4, column=10, value=0)

items = [
    ("Candlestick", "铜制烛台", 1),
    ("Sandal", "凉鞋", 2),
    ("SilverKey", "银钥匙", 3),
]
for i, (name, alias, val) in enumerate(items):
    row = 6 + i
    ws.cell(row=row, column=8, value=name)
    ws.cell(row=row, column=9, value=alias)
    ws.cell(row=row, column=10, value=val)

ws.column_dimensions["A"].width = 12
ws.column_dimensions["H"].width = 20
ws.column_dimensions["I"].width = 15

wb.save("demo_styled/base.xlsx")

# ours: change Candlestick value, add new row
wb2 = openpyxl.load_workbook("demo_styled/base.xlsx")
ws2 = wb2["ItemEnum"]
ws2.cell(row=6, column=10, value=10)
ws2.cell(row=9, column=8, value="VintageWatch")
ws2.cell(row=9, column=9, value="老式手表")
ws2.cell(row=9, column=10, value=4)
wb2.save("demo_styled/ours.xlsx")

# theirs: change SilverKey alias, conflict on Candlestick value
wb3 = openpyxl.load_workbook("demo_styled/base.xlsx")
ws3 = wb3["ItemEnum"]
ws3.cell(row=8, column=9, value="金钥匙")
ws3.cell(row=6, column=10, value=100)
wb3.save("demo_styled/theirs.xlsx")

print("Styled demo files created")
