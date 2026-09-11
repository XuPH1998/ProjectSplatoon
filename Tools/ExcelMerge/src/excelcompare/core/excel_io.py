from __future__ import annotations

import shutil
from copy import copy
from pathlib import Path

import openpyxl

from excelcompare.core.models import (
    CellResolution,
    CellValue,
    DiffKind,
    ResolutionChoice,
    SheetData,
    SheetDiffKind,
    WorkbookData,
    WorkbookDiff,
)


def load_workbook(path: str | Path) -> WorkbookData:
    wb = openpyxl.load_workbook(str(path), data_only=False)
    result = WorkbookData()
    result.sheet_order = wb.sheetnames[:]

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet = SheetData(name=sheet_name)
        sheet.max_row = ws.max_row or 0
        sheet.max_col = ws.max_column or 0

        for row in ws.iter_rows(min_row=1, max_row=sheet.max_row, max_col=sheet.max_col):
            for cell in row:
                if cell.value is None:
                    continue
                cv = CellValue()
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cv.formula = cell.value
                    cv.value = cell.value
                    cv.data_type = "f"
                else:
                    cv.value = cell.value
                    cv.data_type = cell.data_type or "n"
                sheet.cells[(cell.row, cell.column)] = cv

        result.sheets[sheet_name] = sheet

    wb.close()
    return result


def save_workbook(wb_data: WorkbookData, path: str | Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    for sheet_name in wb_data.sheet_order:
        if sheet_name not in wb_data.sheets:
            continue
        sheet = wb_data.sheets[sheet_name]
        ws = wb.create_sheet(title=sheet_name)
        for (row, col), cv in sheet.cells.items():
            if cv.formula:
                ws.cell(row=row, column=col, value=cv.formula)
            else:
                ws.cell(row=row, column=col, value=cv.value)

    wb.save(str(path))
    wb.close()


def save_merged_xlsx(
    ours_path: str | Path,
    theirs_path: str | Path,
    output_path: str | Path,
    diff: WorkbookDiff,
    resolutions: dict[tuple[str, int, int], CellResolution],
    theirs: WorkbookData,
    keep_both_edits: dict | None = None,
) -> None:
    """Save merged result by patching ours xlsx — preserves all styles."""
    shutil.copy2(str(ours_path), str(output_path))
    wb = openpyxl.load_workbook(str(output_path))

    theirs_wb = None
    ours_wb_src = None

    for sheet_name, sheet_diff in diff.sheets.items():
        if sheet_diff.kind == SheetDiffKind.ADDED_THEIRS:
            if theirs_wb is None:
                theirs_wb = openpyxl.load_workbook(str(theirs_path))
            _copy_sheet_with_styles(theirs_wb[sheet_name], wb)
            continue

        if sheet_name not in wb.sheetnames:
            continue

        if sheet_diff.aligned or sheet_diff.col_aligned:
            if theirs_wb is None:
                theirs_wb = openpyxl.load_workbook(str(theirs_path))
            if ours_wb_src is None:
                ours_wb_src = openpyxl.load_workbook(str(ours_path))
            _save_aligned_sheet(
                wb, ours_wb_src, theirs_wb,
                sheet_name, sheet_diff, resolutions, theirs,
                keep_both_edits=keep_both_edits or {},
            )
            continue

        ws = wb[sheet_name]

        for (row, col), cell_diff in sheet_diff.cells.items():
            resolution = resolutions.get((sheet_name, row, col))
            new_value = _resolve_cell_value(cell_diff, resolution, theirs)

            if new_value is _SKIP:
                continue

            need_theirs_style = False
            if resolution and resolution.choice == ResolutionChoice.ACCEPT_THEIRS:
                need_theirs_style = True
            elif not resolution and cell_diff.kind == DiffKind.THEIRS_ONLY:
                need_theirs_style = True

            if need_theirs_style and theirs_wb is None:
                theirs_wb = openpyxl.load_workbook(str(theirs_path))

            cell = ws.cell(row=row, column=col)
            cell.value = new_value

            if need_theirs_style and theirs_wb and sheet_name in theirs_wb.sheetnames:
                src_cell = theirs_wb[sheet_name].cell(row=row, column=col)
                if src_cell.has_style:
                    cell.font = copy(src_cell.font)
                    cell.fill = copy(src_cell.fill)
                    cell.border = copy(src_cell.border)
                    cell.alignment = copy(src_cell.alignment)
                    cell.number_format = src_cell.number_format
                    cell.protection = copy(src_cell.protection)

    if theirs_wb:
        theirs_wb.close()
    if ours_wb_src:
        ours_wb_src.close()

    wb.save(str(output_path))
    wb.close()


_SKIP = object()


def _resolve_cell_value(cell_diff, resolution, theirs: WorkbookData):
    """Determine what value to write. Returns _SKIP if no change needed (keep ours as-is)."""
    if resolution:
        if resolution.choice == ResolutionChoice.ACCEPT_OURS:
            return _SKIP
        elif resolution.choice == ResolutionChoice.ACCEPT_THEIRS:
            theirs_sheet = theirs.sheets.get(resolution.sheet)
            cv = _get_cell(theirs_sheet, (resolution.row, resolution.col))
            return _cv_to_write_value(cv)
        elif resolution.choice == ResolutionChoice.KEEP_BOTH:
            return _SKIP
        elif resolution.choice == ResolutionChoice.DELETE:
            return _SKIP
        else:
            return resolution.manual_value
    else:
        if cell_diff.kind == DiffKind.OURS_ONLY:
            return _SKIP
        elif cell_diff.kind == DiffKind.BOTH_SAME:
            return _SKIP
        elif cell_diff.kind == DiffKind.THEIRS_ONLY:
            return _cv_to_write_value(cell_diff.theirs)
        elif cell_diff.kind == DiffKind.CONFLICT:
            return _SKIP
        return _SKIP


def _cv_to_write_value(cv: CellValue | None):
    if cv is None:
        return None
    if cv.formula:
        return cv.formula
    return cv.value


def _copy_sheet_with_styles(src_ws, dst_wb):
    """Copy a sheet from one workbook to another, preserving styles."""
    dst_ws = dst_wb.create_sheet(title=src_ws.title)
    for row in src_ws.iter_rows():
        for cell in row:
            dst_cell = dst_ws.cell(row=cell.row, column=cell.column, value=cell.value)
            if cell.has_style:
                dst_cell.font = copy(cell.font)
                dst_cell.fill = copy(cell.fill)
                dst_cell.border = copy(cell.border)
                dst_cell.alignment = copy(cell.alignment)
                dst_cell.number_format = cell.number_format
                dst_cell.protection = copy(cell.protection)

    for col_letter, dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[col_letter] = copy(dim)
    for row_idx, dim in src_ws.row_dimensions.items():
        dst_ws.row_dimensions[row_idx] = copy(dim)

    for merge_range in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(merge_range))


def _save_aligned_sheet(
    output_wb, ours_wb, theirs_wb,
    sheet_name, sheet_diff, resolutions, theirs_data,
    keep_both_edits=None,
):
    """Rebuild an aligned sheet: output rows follow the unified row order, columns follow col_mapping."""
    from excelcompare.core.models import DiffKind, ResolutionChoice

    ours_ws = ours_wb[sheet_name]
    theirs_ws = theirs_wb[sheet_name] if sheet_name in theirs_wb.sheetnames else None

    old_ws = output_wb[sheet_name]
    idx = output_wb.sheetnames.index(sheet_name)
    output_wb.remove(old_ws)
    new_ws = output_wb.create_sheet(title=sheet_name, index=idx)

    col_map = sheet_diff.col_mapping
    row_map = sheet_diff.row_mapping
    max_col = sheet_diff.max_col

    def _build_vcol_to_phys(source: str) -> dict[int, int]:
        if not col_map:
            return {}
        key = f"{source}_col"
        return {cm["virtual_col"]: cm[key] for cm in col_map if cm.get(key) is not None}

    ours_v2p = _build_vcol_to_phys("ours")
    theirs_v2p = _build_vcol_to_phys("theirs")

    _col_width_sources: list[tuple[int, object, int]] = []
    if col_map:
        for cm in col_map:
            vcol = cm["virtual_col"]
            ours_pcol = cm.get("ours_col")
            theirs_pcol = cm.get("theirs_col")
            src_ws = None
            src_pcol = None
            if ours_pcol is not None:
                src_ws = ours_ws
                src_pcol = ours_pcol
            elif theirs_pcol is not None and theirs_ws:
                src_ws = theirs_ws
                src_pcol = theirs_pcol
            if src_ws and src_pcol:
                _col_width_sources.append((vcol, src_ws, src_pcol))
    else:
        for col_idx in range(1, (ours_ws.max_column or 0) + 1):
            _col_width_sources.append((col_idx, ours_ws, col_idx))

    def _copy_cell(src_ws, src_row, src_col, dst_row, dst_col):
        if src_ws is None or src_row is None or src_col is None:
            return
        src_cell = src_ws.cell(row=src_row, column=src_col)
        dst_cell = new_ws.cell(row=dst_row, column=dst_col, value=src_cell.value)
        if src_cell.has_style:
            dst_cell.font = copy(src_cell.font)
            dst_cell.fill = copy(src_cell.fill)
            dst_cell.border = copy(src_cell.border)
            dst_cell.alignment = copy(src_cell.alignment)
            dst_cell.number_format = src_cell.number_format
            dst_cell.protection = copy(src_cell.protection)

    def _copy_row_height(src_ws, src_row, dst_row):
        if src_ws is None or src_row is None:
            return
        if src_row in src_ws.row_dimensions:
            src_dim = src_ws.row_dimensions[src_row]
            if src_dim.height is not None:
                new_ws.row_dimensions[dst_row].height = src_dim.height

    out_row = 0

    row_diff_total: dict[int, int] = {}
    col_diff_total: dict[int, int] = {}
    for (r, c) in sheet_diff.cells:
        row_diff_total[r] = row_diff_total.get(r, 0) + 1
        col_diff_total[c] = col_diff_total.get(c, 0) + 1

    row_del_count: dict[int, int] = {}
    col_del_count: dict[int, int] = {}
    for (sn, r, c), res in resolutions.items():
        if sn == sheet_name and res.choice == ResolutionChoice.DELETE:
            row_del_count[r] = row_del_count.get(r, 0) + 1
            col_del_count[c] = col_del_count.get(c, 0) + 1

    new_only_cols: set[int] = set()
    ours_only_cols: set[int] = set()
    theirs_only_cols: set[int] = set()
    if col_map:
        for cm in col_map:
            vcol = cm["virtual_col"]
            has_ours = cm.get("ours_col") is not None
            has_theirs = cm.get("theirs_col") is not None
            if not has_ours or not has_theirs:
                new_only_cols.add(vcol)
            if has_ours and not has_theirs:
                ours_only_cols.add(vcol)
            if has_theirs and not has_ours:
                theirs_only_cols.add(vcol)

    new_only_rows: set[int] = set()
    if row_map:
        for rm in row_map:
            vrow = rm["virtual_row"]
            if rm.get("ours_row") is None or rm.get("theirs_row") is None:
                new_only_rows.add(vrow)

    deleted_cols: set[int] = {
        c for c, cnt in col_del_count.items()
        if cnt >= col_diff_total.get(c, 0) and c in new_only_cols
    }

    col_remap: dict[int, int] = {}
    out_c = 0
    for vcol in range(1, max_col + 1):
        if vcol in deleted_cols:
            continue
        out_c += 1
        col_remap[vcol] = out_c

    for vcol, src_ws_w, src_pcol in _col_width_sources:
        if vcol in deleted_cols:
            continue
        dst_col = col_remap[vcol]
        src_letter = openpyxl.utils.get_column_letter(src_pcol)
        dst_letter = openpyxl.utils.get_column_letter(dst_col)
        if src_letter in src_ws_w.column_dimensions:
            src_dim = src_ws_w.column_dimensions[src_letter]
            new_ws.column_dimensions[dst_letter].width = src_dim.width

    if row_map:
        keep_both_rows: set[int] = set()
        deleted_rows: set[int] = {
            r for r, cnt in row_del_count.items()
            if cnt >= row_diff_total.get(r, 0) and r in new_only_rows
        }
        for (sn, r, c), res in resolutions.items():
            if sn == sheet_name and res.choice == ResolutionChoice.KEEP_BOTH:
                keep_both_rows.add(r)

        for rm in row_map:
            vrow = rm["virtual_row"]
            ours_row = rm.get("ours_row")
            theirs_row = rm.get("theirs_row")
            is_keep_both = vrow in keep_both_rows

            if vrow in deleted_rows:
                continue

            if is_keep_both:
                kb_edits = keep_both_edits or {}
                out_row += 1
                for vcol in range(1, max_col + 1):
                    if vcol in deleted_cols or vcol in theirs_only_cols:
                        continue
                    oc = col_remap[vcol]
                    phys_col = ours_v2p.get(vcol, vcol) if col_map else vcol
                    _copy_cell(ours_ws, ours_row, phys_col, out_row, oc)
                    ek = f"{sheet_name}:{vrow},{vcol}:ours"
                    if ek in kb_edits:
                        new_ws.cell(row=out_row, column=oc).value = kb_edits[ek]
                _copy_row_height(ours_ws, ours_row, out_row)

                out_row += 1
                for vcol in range(1, max_col + 1):
                    if vcol in deleted_cols or vcol in ours_only_cols:
                        continue
                    oc = col_remap[vcol]
                    phys_col = theirs_v2p.get(vcol, vcol) if col_map else vcol
                    _copy_cell(theirs_ws, theirs_row, phys_col, out_row, oc)
                    ek = f"{sheet_name}:{vrow},{vcol}:theirs"
                    if ek in kb_edits:
                        new_ws.cell(row=out_row, column=oc).value = kb_edits[ek]
                _copy_row_height(theirs_ws, theirs_row, out_row)
                continue

            out_row += 1

            src_ws = ours_ws
            src_row = ours_row
            v2p = ours_v2p
            src_is_theirs = False
            if ours_row is None and theirs_row is not None:
                src_ws = theirs_ws
                src_row = theirs_row
                v2p = theirs_v2p
                src_is_theirs = True

            if src_row is not None and src_ws is not None:
                skip_cols = theirs_only_cols if not src_is_theirs else ours_only_cols
                for vcol in range(1, max_col + 1):
                    if vcol in deleted_cols or vcol in skip_cols:
                        continue
                    oc = col_remap[vcol]
                    phys_col = v2p.get(vcol, vcol) if col_map else vcol
                    _copy_cell(src_ws, src_row, phys_col, out_row, oc)

            for vcol in range(1, max_col + 1):
                cell_diff = sheet_diff.cells.get((vrow, vcol))
                if cell_diff is None:
                    continue

                oc = col_remap.get(vcol)
                if oc is None:
                    continue

                resolution = resolutions.get((sheet_name, vrow, vcol))
                if resolution:
                    if resolution.choice == ResolutionChoice.ACCEPT_OURS:
                        if ours_row:
                            phys_col = ours_v2p.get(vcol, vcol) if col_map else vcol
                            _copy_cell(ours_ws, ours_row, phys_col, out_row, oc)
                    elif resolution.choice == ResolutionChoice.ACCEPT_THEIRS:
                        if theirs_row and theirs_ws:
                            phys_col = theirs_v2p.get(vcol, vcol) if col_map else vcol
                            _copy_cell(theirs_ws, theirs_row, phys_col, out_row, oc)
                    elif resolution.choice == ResolutionChoice.DELETE:
                        pass
                    elif resolution.choice == ResolutionChoice.MANUAL:
                        style_ws = ours_ws if ours_row else theirs_ws
                        style_row = ours_row if ours_row else theirs_row
                        style_v2p = ours_v2p if ours_row else theirs_v2p
                        if style_ws and style_row:
                            phys_col = style_v2p.get(vcol, vcol) if col_map else vcol
                            _copy_cell(style_ws, style_row, phys_col, out_row, oc)
                        new_ws.cell(row=out_row, column=oc).value = resolution.manual_value
                elif cell_diff.kind == DiffKind.THEIRS_ONLY:
                    if theirs_row and theirs_ws:
                        phys_col = theirs_v2p.get(vcol, vcol) if col_map else vcol
                        _copy_cell(theirs_ws, theirs_row, phys_col, out_row, oc)
            _copy_row_height(src_ws, src_row, out_row)
    else:
        deleted_rows_nomap: set[int] = {
            r for r, cnt in row_del_count.items()
            if cnt >= row_diff_total.get(r, 0) and r in new_only_rows
        }

        max_row = sheet_diff.max_row
        for row in range(1, max_row + 1):
            if row in deleted_rows_nomap:
                continue
            out_row += 1
            for vcol in range(1, max_col + 1):
                if vcol in deleted_cols:
                    continue
                oc = col_remap[vcol]
                phys_col = ours_v2p.get(vcol, vcol) if col_map else vcol
                _copy_cell(ours_ws, row, phys_col, out_row, oc)

            for vcol in range(1, max_col + 1):
                if vcol in deleted_cols:
                    continue
                oc = col_remap[vcol]
                cell_diff = sheet_diff.cells.get((row, vcol))
                if cell_diff is None:
                    continue
                resolution = resolutions.get((sheet_name, row, vcol))
                if resolution:
                    if resolution.choice == ResolutionChoice.ACCEPT_THEIRS:
                        phys_col = theirs_v2p.get(vcol, vcol) if col_map else vcol
                        if theirs_ws:
                            _copy_cell(theirs_ws, row, phys_col, out_row, oc)
                    elif resolution.choice == ResolutionChoice.DELETE:
                        pass
                    elif resolution.choice == ResolutionChoice.MANUAL:
                        phys_col = ours_v2p.get(vcol, vcol) if col_map else vcol
                        _copy_cell(ours_ws, row, phys_col, out_row, oc)
                        new_ws.cell(row=out_row, column=oc).value = resolution.manual_value
                elif cell_diff.kind == DiffKind.THEIRS_ONLY:
                    phys_col = theirs_v2p.get(vcol, vcol) if col_map else vcol
                    if theirs_ws:
                        _copy_cell(theirs_ws, row, phys_col, out_row, oc)
            _copy_row_height(ours_ws, row, out_row)


def apply_resolutions(
    diff: WorkbookDiff,
    resolutions: dict[tuple[str, int, int], CellResolution],
    base: WorkbookData,
    ours: WorkbookData,
    theirs: WorkbookData,
) -> WorkbookData:
    merged = WorkbookData()
    merged.sheet_order = diff.sheet_order[:]

    for sheet_name in diff.sheet_order:
        sheet_diff = diff.sheets.get(sheet_name)
        if not sheet_diff:
            continue

        ours_sheet = ours.sheets.get(sheet_name)
        theirs_sheet = theirs.sheets.get(sheet_name)
        base_sheet = base.sheets.get(sheet_name)

        merged_sheet = SheetData(name=sheet_name)
        source = ours_sheet or theirs_sheet or base_sheet
        if source:
            merged_sheet.max_row = source.max_row
            merged_sheet.max_col = source.max_col

        all_keys: set[tuple[int, int]] = set()
        for s in (base_sheet, ours_sheet, theirs_sheet):
            if s:
                all_keys.update(s.cells.keys())

        for key in all_keys:
            row, col = key
            cell_diff = sheet_diff.cells.get(key)
            resolution = resolutions.get((sheet_name, row, col))

            if resolution:
                if resolution.choice == ResolutionChoice.DELETE:
                    continue
                if resolution.choice == ResolutionChoice.ACCEPT_OURS:
                    cv = _get_cell(ours_sheet, key)
                elif resolution.choice == ResolutionChoice.ACCEPT_THEIRS:
                    cv = _get_cell(theirs_sheet, key)
                elif resolution.choice == ResolutionChoice.KEEP_BOTH:
                    cv = _get_cell(ours_sheet, key)
                else:
                    cv = CellValue()
                    if resolution.manual_value and resolution.manual_value.startswith("="):
                        cv.formula = resolution.manual_value
                        cv.value = resolution.manual_value
                        cv.data_type = "f"
                    else:
                        cv.value = resolution.manual_value
                        cv.data_type = "s"
            elif cell_diff:
                if cell_diff.kind == DiffKind.OURS_ONLY:
                    cv = _get_cell(ours_sheet, key)
                elif cell_diff.kind == DiffKind.THEIRS_ONLY:
                    cv = _get_cell(theirs_sheet, key)
                elif cell_diff.kind == DiffKind.BOTH_SAME:
                    cv = _get_cell(ours_sheet, key)
                elif cell_diff.kind == DiffKind.CONFLICT:
                    cv = _get_cell(ours_sheet, key)
                else:
                    cv = _get_cell(ours_sheet, key) or _get_cell(base_sheet, key)
            else:
                cv = _get_cell(ours_sheet, key) or _get_cell(base_sheet, key)

            if cv and not cv.is_empty():
                merged_sheet.cells[key] = cv

        if merged_sheet.cells:
            max_r = max(r for r, _ in merged_sheet.cells.keys())
            max_c = max(c for _, c in merged_sheet.cells.keys())
            merged_sheet.max_row = max(merged_sheet.max_row, max_r)
            merged_sheet.max_col = max(merged_sheet.max_col, max_c)

        merged.sheets[sheet_name] = merged_sheet

    return merged


def _get_cell(sheet: SheetData | None, key: tuple[int, int]) -> CellValue | None:
    if sheet is None:
        return None
    return sheet.cells.get(key)
